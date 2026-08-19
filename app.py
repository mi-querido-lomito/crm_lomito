from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import Image as RLImage
from reportlab.graphics.shapes import Drawing
from reportlab.platypus import KeepTogether
from reportlab.lib.units import mm
from reportlab.graphics.barcode import code128
from openpyxl import Workbook
from flask import Flask, render_template, request, redirect, url_for, flash
from flask_mysqldb import MySQL
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from werkzeug.security import check_password_hash
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.platypus import Table
import base64
import io
import json
import os
import textwrap
import uuid
from flask import send_file
from flask import session
import csv
from config import Config
import MySQLdb.cursors
from functools import wraps
from flask import abort
from datetime import datetime, time, timedelta
from werkzeug.utils import secure_filename

def actualizar_citas_perdidas():
    cursor = mysql.connection.cursor()
    cursor.execute("""
        UPDATE citas
        SET estado='Cancelada'
        WHERE estado='Pendiente'
        AND fecha < NOW()
    """)
    mysql.connection.commit()
    cursor.close()


def rol_requerido(roles):
    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            if "rol" not in session:
                abort(403)

            if session["rol"] not in roles:
                abort(403)

            return f(*args, **kwargs)
        return decorated_function
    return decorator

app = Flask(__name__)
app.config.from_object(Config)

schema_checked = False


mysql = MySQL()
mysql.init_app(app)

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "login"

BARCODE_ENTITIES = {
    "productos": "PROD",
    "servicios": "SERV",
    "insumos": "INS",
}

GOOGLE_DRIVE_SCOPES = ["https://www.googleapis.com/auth/drive"]
MASCOTA_FOTO_FORMATOS = {
    "JPEG": ("jpg", "image/jpeg"),
    "PNG": ("png", "image/png"),
    "GIF": ("gif", "image/gif"),
    "WEBP": ("webp", "image/webp"),
}


@app.errorhandler(413)
def archivo_demasiado_grande(error):
    flash("La imagen supera el tamaño máximo permitido")
    return redirect(request.referrer or url_for("dashboard"))


@app.errorhandler(500)
def error_interno(error):
    app.logger.exception("Error interno no controlado")
    flash("Ocurrió un error inesperado. Por favor intenta de nuevo.")
    return redirect(request.referrer or url_for("dashboard"))


def obtener_google_drive_folder_id():
    folder_id = (app.config.get("GOOGLE_DRIVE_FOLDER_ID") or "").strip()
    if not folder_id:
        raise RuntimeError("Configura GOOGLE_DRIVE_FOLDER_ID con el ID de la carpeta de Drive")
    return folder_id


def cargar_google_drive_credentials():
    oauth_refresh_token = (app.config.get("GOOGLE_DRIVE_OAUTH_REFRESH_TOKEN") or "").strip()
    oauth_client_id = (app.config.get("GOOGLE_DRIVE_OAUTH_CLIENT_ID") or "").strip()
    oauth_client_secret = (app.config.get("GOOGLE_DRIVE_OAUTH_CLIENT_SECRET") or "").strip()
    oauth_token_uri = (
        app.config.get("GOOGLE_DRIVE_OAUTH_TOKEN_URI")
        or "https://oauth2.googleapis.com/token"
    ).strip()

    if oauth_refresh_token:
        if not oauth_client_id or not oauth_client_secret:
            raise RuntimeError(
                "Configura GOOGLE_DRIVE_OAUTH_CLIENT_ID y "
                "GOOGLE_DRIVE_OAUTH_CLIENT_SECRET junto con GOOGLE_DRIVE_OAUTH_REFRESH_TOKEN"
            )

        try:
            from google.oauth2.credentials import Credentials
        except ImportError as exc:
            raise RuntimeError("Faltan dependencias de Google Drive; ejecuta pip install -r requirements.txt") from exc

        return Credentials(
            token=None,
            refresh_token=oauth_refresh_token,
            token_uri=oauth_token_uri,
            client_id=oauth_client_id,
            client_secret=oauth_client_secret,
            scopes=GOOGLE_DRIVE_SCOPES
        )

    try:
        from google.oauth2 import service_account
    except ImportError as exc:
        raise RuntimeError("Faltan dependencias de Google Drive; ejecuta pip install -r requirements.txt") from exc

    credenciales_b64 = (app.config.get("GOOGLE_DRIVE_SERVICE_ACCOUNT_JSON_B64") or "").strip()
    credenciales_json = (app.config.get("GOOGLE_DRIVE_SERVICE_ACCOUNT_JSON") or "").strip()
    credenciales_file = (
        app.config.get("GOOGLE_DRIVE_SERVICE_ACCOUNT_FILE")
        or os.getenv("GOOGLE_APPLICATION_CREDENTIALS")
        or ""
    ).strip()

    try:
        if credenciales_b64:
            credenciales_json = base64.b64decode(credenciales_b64).decode("utf-8")

        if credenciales_json:
            return service_account.Credentials.from_service_account_info(
                json.loads(credenciales_json),
                scopes=GOOGLE_DRIVE_SCOPES
            )

        if credenciales_file:
            return service_account.Credentials.from_service_account_file(
                credenciales_file,
                scopes=GOOGLE_DRIVE_SCOPES
            )
    except Exception as exc:
        raise RuntimeError("Las credenciales de Google Drive no se pudieron leer o no tienen un JSON válido") from exc

    raise RuntimeError(
        "Configura GOOGLE_DRIVE_SERVICE_ACCOUNT_JSON_B64, "
        "GOOGLE_DRIVE_SERVICE_ACCOUNT_JSON o GOOGLE_DRIVE_SERVICE_ACCOUNT_FILE"
    )


def crear_google_drive_service():
    try:
        from googleapiclient.discovery import build
    except ImportError as exc:
        raise RuntimeError("Faltan dependencias de Google Drive; ejecuta pip install -r requirements.txt") from exc

    return build(
        "drive",
        "v3",
        credentials=cargar_google_drive_credentials(),
        cache_discovery=False
    )


def validar_foto_imagen(archivo):
    if not archivo or not archivo.filename:
        return None

    contenido = archivo.read()
    max_bytes = int(app.config.get("MASCOTAS_FOTO_MAX_BYTES") or (8 * 1024 * 1024))

    if not contenido:
        raise ValueError("La foto seleccionada está vacía")

    if len(contenido) > max_bytes:
        max_mb = max_bytes // (1024 * 1024)
        raise ValueError(f"La foto no debe superar {max_mb} MB")

    try:
        from PIL import Image, UnidentifiedImageError

        with Image.open(io.BytesIO(contenido)) as imagen:
            formato = imagen.format
            imagen.verify()
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        raise ValueError("La foto debe ser una imagen JPG, PNG, GIF o WebP") from exc

    if formato not in MASCOTA_FOTO_FORMATOS:
        raise ValueError("La foto debe ser una imagen JPG, PNG, GIF o WebP")

    extension, mime_type = MASCOTA_FOTO_FORMATOS[formato]

    return {
        "contenido": contenido,
        "extension": extension,
        "mime_type": mime_type,
    }


def subir_foto_drive(archivo, prefijo, entidad_id, nombre):
    foto = validar_foto_imagen(archivo)
    if not foto:
        return None

    try:
        from googleapiclient.http import MediaIoBaseUpload
    except ImportError as exc:
        raise RuntimeError("Faltan dependencias de Google Drive; ejecuta pip install -r requirements.txt") from exc

    service = crear_google_drive_service()
    nombre_limpio = secure_filename(nombre or prefijo) or prefijo
    nombre_drive = f"{prefijo}_{entidad_id}_{nombre_limpio}_{uuid.uuid4().hex}.{foto['extension']}"

    media = MediaIoBaseUpload(
        io.BytesIO(foto["contenido"]),
        mimetype=foto["mime_type"],
        resumable=True
    )
    metadata = {
        "name": nombre_drive,
        "parents": [obtener_google_drive_folder_id()],
    }

    try:
        archivo_drive = service.files().create(
            body=metadata,
            media_body=media,
            fields="id,name,mimeType,webViewLink",
            supportsAllDrives=True
        ).execute()
    except Exception as exc:
        raise RuntimeError(f"No se pudo subir la foto a Google Drive: {exc}") from exc

    if not archivo_drive.get("id"):
        raise RuntimeError("Google Drive no devolvió el ID del archivo subido")

    return {
        "foto_drive_id": archivo_drive.get("id"),
        "foto_drive_nombre": archivo_drive.get("name") or nombre_drive,
        "foto_mime_type": archivo_drive.get("mimeType") or foto["mime_type"],
        "foto_drive_web_view_url": archivo_drive.get("webViewLink"),
    }


def eliminar_foto_drive(file_id):
    if not file_id:
        return

    try:
        crear_google_drive_service().files().delete(
            fileId=file_id,
            supportsAllDrives=True
        ).execute()
    except Exception as exc:
        status = getattr(getattr(exc, "resp", None), "status", None)
        if status == 404:
            return
        raise RuntimeError(f"No se pudo eliminar la foto anterior de Google Drive: {exc}") from exc


def eliminar_foto_drive_silencioso(file_id):
    try:
        eliminar_foto_drive(file_id)
    except RuntimeError:
        pass


def reconectar_mysql_si_hace_falta():
    # Las llamadas a Google Drive pueden tardar varios segundos; si la
    # conexion a MySQL queda inactiva ese tiempo, el host remoto puede
    # cerrarla y la siguiente consulta fallaria con "server has gone away".
    try:
        mysql.connection.ping(True)
    except Exception:
        pass


def descargar_foto_drive(file_id):
    try:
        from googleapiclient.http import MediaIoBaseDownload
    except ImportError as exc:
        raise RuntimeError("Faltan dependencias de Google Drive; ejecuta pip install -r requirements.txt") from exc

    buffer = io.BytesIO()
    try:
        drive_request = crear_google_drive_service().files().get_media(
            fileId=file_id,
            supportsAllDrives=True
        )
        downloader = MediaIoBaseDownload(buffer, drive_request)
        terminado = False

        while not terminado:
            _, terminado = downloader.next_chunk()
    except Exception as exc:
        raise RuntimeError(f"No se pudo descargar la foto desde Google Drive: {exc}") from exc

    buffer.seek(0)
    return buffer


def generar_codigo_barra(prefijo):
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    codigo = generar_codigo_barra_cursor(cursor, prefijo)
    cursor.close()
    return codigo


def generar_codigo_barra_cursor(cursor, prefijo):
    cursor.execute("""
        SELECT MAX(CAST(SUBSTRING_INDEX(codigo_barras, '-', -1) AS UNSIGNED)) AS ultimo
        FROM (
            SELECT codigo_barras FROM productos
            UNION ALL SELECT codigo_barras FROM servicios
            UNION ALL SELECT codigo_barras FROM insumos
        ) AS codigos
        WHERE codigo_barras REGEXP %s
    """, (f"^{prefijo}-[0-9]+$",))

    row = cursor.fetchone()

    ultimo = int(row["ultimo"]) if row and row.get("ultimo") else 1000
    return f"{prefijo}-{ultimo + 1}"


def asegurar_codigo_barra_actual(tabla, registro_id):
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute(f"SELECT codigo_barras FROM {tabla} WHERE id=%s", (registro_id,))
    registro = cursor.fetchone()

    if not registro:
        cursor.close()
        return None

    codigo = (registro.get("codigo_barras") or "").strip()
    if not codigo:
        codigo = generar_codigo_barra_cursor(cursor, BARCODE_ENTITIES[tabla])
        cursor.execute(
            f"UPDATE {tabla} SET codigo_barras=%s WHERE id=%s",
            (codigo, registro_id)
        )
        mysql.connection.commit()

    cursor.close()
    return codigo


def generar_etiqueta_png(nombre, precio, codigo_barras):
    from PIL import Image, ImageDraw, ImageFont
    import barcode
    from barcode.writer import ImageWriter

    ancho, alto = 480, 320  # ~60mm x 40mm a 8 px/mm (mismo tamaño que la etiqueta HTML)
    lienzo = Image.new("RGB", (ancho, alto), "white")
    dibujo = ImageDraw.Draw(lienzo)

    fuente_nombre = ImageFont.load_default(size=26)
    fuente_precio = ImageFont.load_default(size=22)
    fuente_codigo = ImageFont.load_default(size=16)

    y = 12
    for linea in textwrap.wrap(nombre or "", width=24)[:3]:
        ancho_texto = dibujo.textlength(linea, font=fuente_nombre)
        dibujo.text(((ancho - ancho_texto) / 2, y), linea, fill="black", font=fuente_nombre)
        y += 30

    texto_precio = f"${float(precio or 0):.2f}"
    ancho_precio = dibujo.textlength(texto_precio, font=fuente_precio)
    dibujo.text(((ancho - ancho_precio) / 2, y + 4), texto_precio, fill="black", font=fuente_precio)
    y += 44

    if codigo_barras:
        clase_codigo = barcode.get_barcode_class("code128")
        codigo_obj = clase_codigo(str(codigo_barras), writer=ImageWriter())
        buffer_codigo = io.BytesIO()
        codigo_obj.write(buffer_codigo, options={
            "write_text": False,
            "module_height": 12,
            "quiet_zone": 2,
        })
        buffer_codigo.seek(0)

        imagen_codigo = Image.open(buffer_codigo)
        # El texto del codigo y el margen inferior necesitan ~26px propios;
        # la escala respeta tanto el ancho como el alto restante del lienzo.
        ancho_max = ancho - 40
        alto_max = max(alto - y - 26, 10)
        escala = min(ancho_max / imagen_codigo.width, alto_max / imagen_codigo.height)
        nuevo_tamano = (int(imagen_codigo.width * escala), int(imagen_codigo.height * escala))
        imagen_codigo = imagen_codigo.resize(nuevo_tamano)
        lienzo.paste(imagen_codigo, ((ancho - nuevo_tamano[0]) // 2, y))
        y += nuevo_tamano[1] + 6

        ancho_codigo_txt = dibujo.textlength(str(codigo_barras), font=fuente_codigo)
        dibujo.text(((ancho - ancho_codigo_txt) / 2, y), str(codigo_barras), fill="black", font=fuente_codigo)

    salida = io.BytesIO()
    lienzo.save(salida, format="PNG")
    salida.seek(0)
    return salida


def rellenar_codigos_barras_faltantes(tabla):
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute(f"""
        SELECT id
        FROM {tabla}
        WHERE codigo_barras IS NULL OR TRIM(codigo_barras) = ''
        ORDER BY id
    """)
    registros = cursor.fetchall()
    for registro in registros:
        codigo = generar_codigo_barra_cursor(cursor, BARCODE_ENTITIES[tabla])
        cursor.execute(
            f"UPDATE {tabla} SET codigo_barras=%s WHERE id=%s",
            (codigo, registro["id"])
        )

    if registros:
        mysql.connection.commit()

    cursor.close()


def columna_existe(cursor, tabla, columna):
    cursor.execute(f"SHOW COLUMNS FROM {tabla} LIKE %s", (columna,))
    return cursor.fetchone() is not None


def asegurar_columna(cursor, tabla, columna, definicion):
    if not columna_existe(cursor, tabla, columna):
        cursor.execute(f"ALTER TABLE {tabla} ADD COLUMN {columna} {definicion}")


def indice_existe(cursor, tabla, indice):
    cursor.execute(f"SHOW INDEX FROM {tabla} WHERE Key_name=%s", (indice,))
    return cursor.fetchone() is not None


def asegurar_indice(cursor, tabla, indice, columnas):
    if not indice_existe(cursor, tabla, indice):
        cursor.execute(f"ALTER TABLE {tabla} ADD INDEX {indice} ({columnas})")


def asegurar_tabla_proveedores(cursor):
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS proveedores (
            id INT NOT NULL AUTO_INCREMENT,
            nombre VARCHAR(150) NOT NULL,
            direccion_contacto TEXT NULL,
            telefono VARCHAR(50) NULL,
            observaciones TEXT NULL,
            PRIMARY KEY (id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb3 COLLATE=utf8mb3_spanish2_ci
    """)


def asegurar_schema_crm():
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    asegurar_tabla_proveedores(cursor)

    columnas = {
        "clientes": {
            "direccion": "TEXT NULL",
        },
        "productos": {
            "proveedor_id": "INT NULL",
            "foto_drive_id": "VARCHAR(255) NULL",
            "foto_drive_nombre": "VARCHAR(255) NULL",
            "foto_mime_type": "VARCHAR(100) NULL",
            "foto_drive_web_view_url": "TEXT NULL",
            "foto_actualizada": "DATETIME NULL",
        },
        "insumos": {
            "proveedor_id": "INT NULL",
            "foto_drive_id": "VARCHAR(255) NULL",
            "foto_drive_nombre": "VARCHAR(255) NULL",
            "foto_mime_type": "VARCHAR(100) NULL",
            "foto_drive_web_view_url": "TEXT NULL",
            "foto_actualizada": "DATETIME NULL",
        },
        "mascotas": {
            "raza": "VARCHAR(100) NULL",
            "edad": "INT NULL",
            "fecha_especial": "DATE NULL",
            "alergias": "TEXT NULL",
            "genero": "VARCHAR(20) NULL",
            "talla": "VARCHAR(50) NULL",
            "color": "VARCHAR(80) NULL",
            "notas": "TEXT NULL",
            "foto_drive_id": "VARCHAR(255) NULL",
            "foto_drive_nombre": "VARCHAR(255) NULL",
            "foto_mime_type": "VARCHAR(100) NULL",
            "foto_drive_web_view_url": "TEXT NULL",
            "foto_actualizada": "DATETIME NULL",
        },
        "citas": {
            "servicio_id": "INT NULL",
            "observaciones": "TEXT NULL",
            "tiempo_servicio_min": "INT NULL",
        },
        "ventas": {
            "mascota_id": "INT NULL",
            "cita_id": "INT NULL",
        },
        "detalle_venta": {
            "mascota_id": "INT NULL",
        },
    }

    for tabla, definiciones in columnas.items():
        for columna, definicion in definiciones.items():
            asegurar_columna(cursor, tabla, columna, definicion)

    # Los codigos de barras se buscan por igualdad en cada venta/entrada
    # de inventario; sin indice, cada escaneo hace un table scan completo.
    indices = {
        "productos": "idx_productos_codigo_barras",
        "servicios": "idx_servicios_codigo_barras",
        "insumos": "idx_insumos_codigo_barras",
    }

    for tabla, nombre_indice in indices.items():
        asegurar_indice(cursor, tabla, nombre_indice, "codigo_barras")

    mysql.connection.commit()
    cursor.close()


def estimar_tiempo_mascota(cursor, mascota_id):
    if not mascota_id:
        return 60

    cursor.execute("""
        SELECT AVG(tiempo_servicio_min) AS promedio
        FROM citas
        WHERE mascota_id=%s
        AND estado='Completada'
        AND tiempo_servicio_min IS NOT NULL
        AND tiempo_servicio_min > 0
    """, (mascota_id,))
    promedio = cursor.fetchone()

    if promedio and promedio.get("promedio"):
        return int(round(float(promedio["promedio"])))

    return 60


def promedios_tiempo_mascotas(cursor, mascota_ids):
    ids = list({mid for mid in mascota_ids if mid})
    if not ids:
        return {}

    placeholders = ", ".join(["%s"] * len(ids))
    cursor.execute(f"""
        SELECT mascota_id, AVG(tiempo_servicio_min) AS promedio
        FROM citas
        WHERE mascota_id IN ({placeholders})
        AND estado='Completada'
        AND tiempo_servicio_min IS NOT NULL
        AND tiempo_servicio_min > 0
        GROUP BY mascota_id
    """, tuple(ids))

    return {
        row["mascota_id"]: int(round(float(row["promedio"])))
        for row in cursor.fetchall()
        if row.get("promedio")
    }


def resumen_agenda_dia(cursor, fecha, mascota_id=None):
    inicio = datetime.combine(fecha.date(), time(10, 0))
    fin = datetime.combine(fecha.date(), time(17, 0))
    jornada_min = int((fin - inicio).total_seconds() / 60)

    cursor.execute("""
        SELECT mascota_id
        FROM citas
        WHERE DATE(fecha)=DATE(%s)
        AND estado='Pendiente'
    """, (fecha,))
    citas_dia = cursor.fetchall()

    ids_a_consultar = [cita.get("mascota_id") for cita in citas_dia]
    if mascota_id:
        ids_a_consultar.append(mascota_id)

    promedios = promedios_tiempo_mascotas(cursor, ids_a_consultar)
    nuevo_min = promedios.get(mascota_id, 60) if mascota_id else 60
    ocupado = sum(promedios.get(cita.get("mascota_id"), 60) for cita in citas_dia)

    return {
        "inicio": inicio,
        "fin": fin,
        "jornada_min": jornada_min,
        "ocupado_min": ocupado,
        "disponible_min": max(jornada_min - ocupado, 0),
        "nuevo_min": nuevo_min,
        "cabe": ocupado + nuevo_min <= jornada_min and inicio <= fecha < fin,
    }


def cerrar_cajas_vencidas(cursor):
    cursor.execute("""
        SELECT *
        FROM caja
        WHERE estado='abierta'
        AND DATE(fecha_apertura) < CURDATE()
        ORDER BY fecha_apertura ASC
    """)
    cajas_vencidas = cursor.fetchall()

    for caja in cajas_vencidas:
        fecha_apertura = caja["fecha_apertura"]
        cierre_dia = datetime.combine(fecha_apertura.date(), time(23, 59, 59))

        cursor.execute("""
            SELECT IFNULL(SUM(total),0) as total_ventas
            FROM ventas
            WHERE estado='activa'
            AND fecha >= %s
            AND fecha <= %s
        """, (fecha_apertura, cierre_dia))
        resultado = cursor.fetchone()
        total_ventas = float(resultado["total_ventas"] or 0) if resultado else 0
        monto_final = float(caja["monto_inicial"] or 0) + total_ventas

        cursor.execute("""
            UPDATE caja
            SET fecha_cierre=%s,
                monto_final=%s,
                diferencia=0,
                estado='cerrada'
            WHERE id=%s
        """, (cierre_dia, monto_final, caja["id"]))


@app.before_request
def verificar_schema_crm():
    global schema_checked
    if schema_checked or request.endpoint in ("static", "login", "home"):
        return

    asegurar_schema_crm()
    schema_checked = True


@app.before_request
def verificar_caja_diaria():
    if request.endpoint in (None, "static", "login", "home", "logout"):
        return

    if not current_user.is_authenticated:
        return

    hoy = datetime.now().date().isoformat()

    if session.get("caja_verificada") == hoy:
        return

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cerrar_cajas_vencidas(cursor)
    mysql.connection.commit()

    cursor.execute("""
        SELECT id
        FROM caja
        WHERE estado='abierta'
        AND DATE(fecha_apertura)=CURDATE()
        LIMIT 1
    """)
    caja_hoy = cursor.fetchone()
    cursor.close()

    if caja_hoy:
        session["caja_verificada"] = hoy
        return

    if request.endpoint == "abrir_caja":
        return

    if datetime.now().time() >= time(9, 30):
        return redirect(url_for("abrir_caja", requerida=1))

    # No se cachea aqui: antes de las 9:30 el estado puede cambiar a
    # "requiere caja" en la siguiente peticion sin que el usuario refresque.

# -------------------
# USER MODEL
# -------------------

class User(UserMixin):
    def __init__(self, user_data):
        self.id = str(user_data["id"])
        self.username = user_data["username"]

@login_manager.user_loader
def load_user(user_id):
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("SELECT * FROM usuarios WHERE id = %s", (user_id,))
    user = cursor.fetchone()
    cursor.close()
    if user:
        return User(user)
    return None


# -------------------
# LOGIN
# -------------------
@app.route("/")
def home():
    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form["username"]
        password = request.form["password"]

        cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
        cursor.execute("SELECT * FROM usuarios WHERE username=%s", (username,))
        user = cursor.fetchone()
        cursor.close()

        password_hash = user.get("PASSWORD") if user else None

        if password_hash and check_password_hash(password_hash, password):
            login_user(User(user))
            session["rol"] = user["rol"]
            return redirect(url_for("dashboard"))
        else:
            flash("Credenciales incorrectas")

    return render_template("login.html")

@app.route("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("login"))


# -------------------
# DASHBOARD
# -------------------
@app.route("/dashboard")
@login_required
def dashboard():

    actualizar_citas_perdidas()

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    rol = session.get("rol")

    # ==========================
    # MÉTRICAS SOLO ADMIN
    # ==========================
    if rol == "admin":

        # Conteos y venta de hoy en una sola consulta (evita 4 viajes
        # redondos separados a la base de datos remota)
        cursor.execute("""
            SELECT
                (SELECT COUNT(*) FROM clientes) AS clientes,
                (SELECT COUNT(*) FROM mascotas) AS mascotas,
                (SELECT COUNT(*) FROM productos) AS productos,
                (SELECT IFNULL(SUM(total),0) FROM ventas
                    WHERE DATE(fecha)=CURDATE() AND estado='activa') AS total_hoy
        """)
        resumen_admin = cursor.fetchone()
        clientes = resumen_admin["clientes"]
        mascotas = resumen_admin["mascotas"]
        productos = resumen_admin["productos"]
        total_hoy = resumen_admin["total_hoy"]

        # Ventas últimos 7 días
        cursor.execute("""
            SELECT DATE(fecha) as dia, SUM(total) as total
            FROM ventas
            WHERE fecha >= CURDATE() - INTERVAL 7 DAY
            AND estado='activa'
            GROUP BY DATE(fecha)
            ORDER BY dia
        """)
        ventas_7_dias = cursor.fetchall()

        # Inventario bajo
        cursor.execute("SELECT nombre, stock FROM productos WHERE stock <= 2")
        bajo_stock = cursor.fetchall()

    else:
        # Empleado no necesita estas métricas
        clientes = None
        mascotas = None
        productos = None
        total_hoy = None
        ventas_7_dias = []
        bajo_stock = []

    # ==========================
    # MÉTRICAS GENERALES (TODOS)
    # ==========================

    # Citas de hoy
    cursor.execute("""
        SELECT estado, COUNT(*) as total
        FROM citas
        WHERE DATE(fecha) = CURDATE()
        GROUP BY estado
    """)
    resumen_citas = cursor.fetchall()

    # Estado de caja
    cursor.execute("""
        SELECT *
        FROM caja
        WHERE estado='abierta'
        AND DATE(fecha_apertura)=CURDATE()
        ORDER BY fecha_apertura DESC
        LIMIT 1
    """)
    caja_abierta = cursor.fetchone()

    # Citas próximas
    cursor.execute("""
        SELECT c.id,
               c.fecha,
               m.nombre as mascota,
               s.nombre as servicio
        FROM citas c
        JOIN mascotas m ON c.mascota_id = m.id
        LEFT JOIN servicios s ON c.servicio_id = s.id
        WHERE c.estado='Pendiente'
        AND c.fecha >= CURDATE()
        AND c.fecha < DATE_ADD(CURDATE(), INTERVAL 7 DAY)
        ORDER BY c.fecha ASC
    """)
    citas_proximas = cursor.fetchall()

    cursor.close()

    return render_template(
        "dashboard.html",
        clientes=clientes,
        mascotas=mascotas,
        productos=productos,
        total_hoy=total_hoy,
        resumen_citas=resumen_citas,
        ventas_7_dias=ventas_7_dias,
        bajo_stock=bajo_stock,
        caja_abierta=caja_abierta,
        citas_proximas=citas_proximas
    )


# -------------------
# CLIENTES
# -------------------

@app.route("/clientes", methods=["GET", "POST"])
@login_required
def clientes():

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    if request.method == "POST":

        nombre = request.form.get("nombre")
        telefono = request.form.get("telefono")
        direccion = request.form.get("direccion")

        if not nombre:
            flash("El nombre es obligatorio")
            return redirect(url_for("clientes"))

        cursor.execute("""
            INSERT INTO clientes (nombre, telefono, direccion)
            VALUES (%s, %s, %s)
        """, (nombre, telefono, direccion))

        mysql.connection.commit()
        flash("Cliente agregado correctamente")

        return redirect(url_for("clientes"))

    cursor.execute("""
        SELECT clientes.*,
               GROUP_CONCAT(mascotas.nombre ORDER BY mascotas.nombre SEPARATOR ', ') AS mascotas_nombres
        FROM clientes
        LEFT JOIN mascotas ON mascotas.cliente_id = clientes.id
        GROUP BY clientes.id
        ORDER BY clientes.id DESC
    """)
    data = cursor.fetchall()
    cursor.close()

    return render_template("clientes.html", clientes=data)


# -------------------
# EDITAR CLIENTE
# -------------------

@app.route("/editar_cliente/<int:id>", methods=["GET", "POST"])
@login_required
def editar_cliente(id):

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cursor.execute("SELECT * FROM clientes WHERE id=%s", (id,))
    cliente = cursor.fetchone()

    if not cliente:
        cursor.close()
        flash("Cliente no encontrado")
        return redirect(url_for("clientes"))

    if request.method == "POST":

        nombre = request.form.get("nombre")
        telefono = request.form.get("telefono")
        direccion = request.form.get("direccion")

        cursor.execute("""
            UPDATE clientes
            SET nombre=%s,
                telefono=%s,
                direccion=%s
            WHERE id=%s
        """, (nombre, telefono, direccion, id))

        mysql.connection.commit()
        cursor.close()

        flash("Cliente actualizado correctamente")
        return redirect(url_for("clientes"))

    cursor.close()
    return render_template("editar_cliente.html", cliente=cliente)


# -------------------
# ELIMINAR CLIENTE (SOLO ADMIN)
# -------------------

@app.route("/eliminar_cliente/<int:id>")
@login_required
@rol_requerido(["admin"])
def eliminar_cliente(id):

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    # Verificar si tiene mascotas asociadas
    cursor.execute("SELECT COUNT(*) as total FROM mascotas WHERE cliente_id=%s", (id,))
    mascotas = cursor.fetchone()["total"]

    if mascotas > 0:
        cursor.close()
        flash("No puedes eliminar un cliente con mascotas registradas")
        return redirect(url_for("clientes"))

    cursor.execute("DELETE FROM clientes WHERE id=%s", (id,))
    mysql.connection.commit()
    cursor.close()

    flash("Cliente eliminado correctamente")
    return redirect(url_for("clientes"))


# -------------------
# EXPORTAR CLIENTES A PDF (SOLO ADMIN)
# -------------------

@app.route("/clientes/pdf")
@login_required
@rol_requerido(["admin"])
def export_clientes_pdf():

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer)
    elements = []
    styles = getSampleStyleSheet()

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("SELECT nombre, telefono, direccion FROM clientes")
    data = cursor.fetchall()
    cursor.close()

    table_data = [["Nombre", "Teléfono", "Dirección"]]

    for c in data:
        table_data.append([
            c["nombre"],
            c["telefono"] if c["telefono"] else "",
            c["direccion"] if c["direccion"] else ""
        ])

    table = Table(table_data)
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="clientes.pdf",
        mimetype="application/pdf"
    )

# -------------------
# MASCOTAS
# -------------------

@app.route("/mascotas", methods=["GET", "POST"])
@login_required
def mascotas():
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    if request.method == "POST":
        try:
            nombre = request.form["nombre"]
            especie = request.form["especie"]
            cliente_id = request.form.get("cliente_id")
            nuevo_cliente = request.form.get("nuevo_cliente")

            if nuevo_cliente == "1":
                cliente_nombre = request.form.get("cliente_nombre")
                cliente_telefono = request.form.get("cliente_telefono")
                cliente_direccion = request.form.get("cliente_direccion")

                if not cliente_nombre:
                    flash("El nombre del nuevo cliente es obligatorio")
                    cursor.close()
                    return redirect(url_for("mascotas"))
            elif not cliente_id:
                flash("Selecciona un dueño o registra un cliente nuevo")
                cursor.close()
                return redirect(url_for("mascotas"))

            # La foto se sube antes de escribir en la BD: si Google Drive
            # tarda y la conexion a MySQL se cae mientras tanto, no queda
            # un cliente/mascota a medio insertar.
            foto_drive = subir_foto_drive(request.files.get("foto"), "mascota", "nueva", nombre)
            reconectar_mysql_si_hace_falta()

            if nuevo_cliente == "1":
                cursor.execute("""
                    INSERT INTO clientes (nombre, telefono, direccion)
                    VALUES (%s, %s, %s)
                """, (cliente_nombre, cliente_telefono, cliente_direccion))
                cliente_id = cursor.lastrowid

            datos = (
                nombre,
                especie,
                cliente_id,
                request.form.get("raza"),
                request.form.get("edad") or None,
                request.form.get("fecha_especial") or None,
                request.form.get("alergias"),
                request.form.get("genero"),
                request.form.get("talla"),
                request.form.get("color"),
                request.form.get("notas"),
            )

            cursor.execute("""
                INSERT INTO mascotas
                (nombre, especie, cliente_id, raza, edad, fecha_especial, alergias, genero, talla, color, notas)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """, datos)
            mascota_id = cursor.lastrowid

            if foto_drive:
                cursor.execute("""
                    UPDATE mascotas
                    SET foto_drive_id=%s,
                        foto_drive_nombre=%s,
                        foto_mime_type=%s,
                        foto_drive_web_view_url=%s,
                        foto_actualizada=NOW()
                    WHERE id=%s
                """, (
                    foto_drive["foto_drive_id"],
                    foto_drive["foto_drive_nombre"],
                    foto_drive["foto_mime_type"],
                    foto_drive["foto_drive_web_view_url"],
                    mascota_id
                ))

            mysql.connection.commit()
            flash("Mascota registrada correctamente")
        except (ValueError, RuntimeError, MySQLdb.OperationalError) as error:
            mysql.connection.rollback()
            flash(str(error))

        cursor.close()
        return redirect(url_for("mascotas"))

    cursor.execute("""
        SELECT mascotas.*, clientes.nombre as dueno
        FROM mascotas
        JOIN clientes ON mascotas.cliente_id = clientes.id
        ORDER BY mascotas.id DESC
    """)
    data = cursor.fetchall()

    cursor.execute("SELECT * FROM clientes")
    clientes = cursor.fetchall()

    cursor.close()

    return render_template("mascotas.html", mascotas=data, clientes=clientes)


# -------------------
# EDITAR MASCOTA
@app.route("/editar_mascota/<int:id>", methods=["GET", "POST"])
@login_required
def editar_mascota(id):

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cursor.execute("SELECT * FROM mascotas WHERE id=%s", (id,))
    mascota = cursor.fetchone()

    if not mascota:
        cursor.close()
        flash("Mascota no encontrada")
        return redirect(url_for("mascotas"))

    # necesitamos lista de clientes para el select
    cursor.execute("SELECT * FROM clientes")
    clientes = cursor.fetchall()

    if request.method == "POST":

        nombre = request.form["nombre"]
        especie = request.form["especie"]
        cliente_id = request.form["cliente_id"]
        raza = request.form.get("raza")
        edad = request.form.get("edad") or None
        fecha_especial = request.form.get("fecha_especial") or None
        alergias = request.form.get("alergias")
        genero = request.form.get("genero")
        talla = request.form.get("talla")
        color = request.form.get("color")
        notas = request.form.get("notas")
        foto_archivo = request.files.get("foto")
        reemplaza_foto = foto_archivo and foto_archivo.filename
        quitar_foto = request.form.get("quitar_foto") == "1" and not reemplaza_foto
        nueva_foto_drive = None

        try:
            if reemplaza_foto:
                nueva_foto_drive = subir_foto_drive(foto_archivo, "mascota", id, nombre)
                reconectar_mysql_si_hace_falta()

            cursor.execute("""
                UPDATE mascotas
                SET nombre=%s,
                    especie=%s,
                    cliente_id=%s,
                    raza=%s,
                    edad=%s,
                    fecha_especial=%s,
                    alergias=%s,
                    genero=%s,
                    talla=%s,
                    color=%s,
                    notas=%s
                WHERE id=%s
            """, (nombre, especie, cliente_id, raza, edad, fecha_especial, alergias, genero, talla, color, notas, id))

            if nueva_foto_drive:
                cursor.execute("""
                    UPDATE mascotas
                    SET foto_drive_id=%s,
                        foto_drive_nombre=%s,
                        foto_mime_type=%s,
                        foto_drive_web_view_url=%s,
                        foto_actualizada=NOW()
                    WHERE id=%s
                """, (
                    nueva_foto_drive["foto_drive_id"],
                    nueva_foto_drive["foto_drive_nombre"],
                    nueva_foto_drive["foto_mime_type"],
                    nueva_foto_drive["foto_drive_web_view_url"],
                    id
                ))
            elif quitar_foto:
                cursor.execute("""
                    UPDATE mascotas
                    SET foto_drive_id=NULL,
                        foto_drive_nombre=NULL,
                        foto_mime_type=NULL,
                        foto_drive_web_view_url=NULL,
                        foto_actualizada=NULL
                    WHERE id=%s
                """, (id,))

            mysql.connection.commit()
        except (ValueError, RuntimeError, MySQLdb.OperationalError) as error:
            mysql.connection.rollback()
            if nueva_foto_drive:
                eliminar_foto_drive_silencioso(nueva_foto_drive["foto_drive_id"])
            flash(str(error))
            cursor.close()
            return redirect(url_for("editar_mascota", id=id))

        if nueva_foto_drive or quitar_foto:
            eliminar_foto_drive_silencioso(mascota.get("foto_drive_id"))

        cursor.close()
        flash("Mascota actualizada correctamente")
        return redirect(url_for("mascotas"))

    cursor.close()
    return render_template("editar_mascota.html",
                           mascota=mascota,
                           clientes=clientes)


@app.route("/mascota/<int:id>/foto")
@login_required
def foto_mascota(id):
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("""
        SELECT foto_drive_id, foto_drive_nombre, foto_mime_type
        FROM mascotas
        WHERE id=%s
    """, (id,))
    mascota = cursor.fetchone()
    cursor.close()

    if not mascota or not mascota.get("foto_drive_id"):
        abort(404)

    try:
        archivo = descargar_foto_drive(mascota["foto_drive_id"])
    except RuntimeError:
        abort(404)

    return send_file(
        archivo,
        mimetype=mascota.get("foto_mime_type") or "image/jpeg",
        download_name=mascota.get("foto_drive_nombre") or f"mascota-{id}.jpg",
        max_age=3600
    )


# ELIMINAR MASCOTA
# -------------------
@app.route("/eliminar_mascota/<int:id>")
@login_required
@rol_requerido(["admin"])
def eliminar_mascota(id):

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("SELECT foto_drive_id FROM mascotas WHERE id=%s", (id,))
    mascota = cursor.fetchone()

    cursor.execute("DELETE FROM mascotas WHERE id=%s", (id,))
    mysql.connection.commit()
    cursor.close()

    foto_error = False
    if mascota and mascota.get("foto_drive_id"):
        try:
            eliminar_foto_drive(mascota["foto_drive_id"])
        except RuntimeError:
            foto_error = True

    if foto_error:
        flash("Mascota eliminada, pero no se pudo borrar su foto de Drive")
    else:
        flash("Mascota eliminada correctamente")
    return redirect(url_for("mascotas"))


@app.route("/mascota/<int:id>/historial")
@login_required
def historial_mascota(id):
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cursor.execute("""
        SELECT mascotas.*, clientes.nombre as dueno, clientes.telefono, clientes.direccion
        FROM mascotas
        JOIN clientes ON mascotas.cliente_id = clientes.id
        WHERE mascotas.id=%s
    """, (id,))
    mascota = cursor.fetchone()

    if not mascota:
        cursor.close()
        flash("Mascota no encontrada")
        return redirect(url_for("mascotas"))

    cursor.execute("""
        SELECT c.*, s.nombre as servicio, s.precio
        FROM citas c
        LEFT JOIN servicios s ON c.servicio_id = s.id
        WHERE c.mascota_id=%s
        ORDER BY c.fecha DESC
    """, (id,))
    citas = cursor.fetchall()

    cursor.execute("""
        SELECT v.id as venta_id, v.fecha, v.total, v.metodo_pago,
               d.tipo, d.cantidad, d.precio_unitario, d.subtotal,
               s.nombre as servicio, p.nombre as producto
        FROM ventas v
        JOIN detalle_venta d ON d.venta_id = v.id
        LEFT JOIN servicios s ON d.servicio_id = s.id
        LEFT JOIN productos p ON d.producto_id = p.id
        WHERE (d.mascota_id=%s OR (d.mascota_id IS NULL AND v.mascota_id=%s))
        AND v.estado='activa'
        ORDER BY v.fecha DESC
    """, (id, id))
    ventas = cursor.fetchall()

    cursor.execute("""
        SELECT s.nombre as servicio, AVG(c.tiempo_servicio_min) as promedio_min, COUNT(*) as total
        FROM citas c
        LEFT JOIN servicios s ON c.servicio_id = s.id
        WHERE c.mascota_id=%s
        AND c.estado='Completada'
        AND c.tiempo_servicio_min IS NOT NULL
        GROUP BY c.servicio_id, s.nombre
        ORDER BY s.nombre
    """, (id,))
    promedios = cursor.fetchall()

    cursor.close()

    return render_template(
        "historial_mascota.html",
        mascota=mascota,
        citas=citas,
        ventas=ventas,
        promedios=promedios
    )


# -------------------
# PRODUCTOS
# -------------------

@app.route("/productos", methods=["GET", "POST"])
@login_required
@rol_requerido(["admin"])
def productos():
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    if request.method == "POST":
        nombre = request.form["nombre"]
        stock = request.form["stock"]
        precio = request.form["precio"]
        proveedor_id = request.form.get("proveedor_id") or None

        try:
            # La foto se sube antes de escribir en la BD para no dejar un
            # producto a medio insertar si Drive tarda y la conexion cae.
            foto_drive = subir_foto_drive(request.files.get("foto"), "producto", "nuevo", nombre)
            reconectar_mysql_si_hace_falta()

            codigo = generar_codigo_barra("PROD")

            cursor.execute(
                "INSERT INTO productos (nombre, stock, precio, proveedor_id, codigo_barras) VALUES (%s,%s,%s,%s,%s)",
                (nombre, stock, precio, proveedor_id, codigo)
            )
            producto_id = cursor.lastrowid

            if foto_drive:
                cursor.execute("""
                    UPDATE productos
                    SET foto_drive_id=%s,
                        foto_drive_nombre=%s,
                        foto_mime_type=%s,
                        foto_drive_web_view_url=%s,
                        foto_actualizada=NOW()
                    WHERE id=%s
                """, (
                    foto_drive["foto_drive_id"],
                    foto_drive["foto_drive_nombre"],
                    foto_drive["foto_mime_type"],
                    foto_drive["foto_drive_web_view_url"],
                    producto_id
                ))

            mysql.connection.commit()
        except (ValueError, RuntimeError, MySQLdb.OperationalError) as error:
            mysql.connection.rollback()
            flash(str(error))

    rellenar_codigos_barras_faltantes("productos")
    cursor.execute("""
        SELECT productos.*, proveedores.nombre AS proveedor_nombre
        FROM productos
        LEFT JOIN proveedores ON productos.proveedor_id = proveedores.id
        ORDER BY productos.id DESC
    """)
    data = cursor.fetchall()
    cursor.execute("SELECT * FROM proveedores ORDER BY nombre")
    proveedores = cursor.fetchall()
    cursor.close()

    return render_template("productos.html", productos=data, proveedores=proveedores)

# EDITAR PRODUCTO
# -------------------
@app.route("/editar_producto/<int:id>", methods=["GET", "POST"])
@login_required
@rol_requerido(["admin"])
def editar_producto(id):

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cursor.execute("SELECT * FROM productos WHERE id=%s", (id,))
    producto = cursor.fetchone()

    if request.method == "POST":

        nombre = request.form["nombre"]
        precio = request.form["precio"]
        stock = request.form["stock"]
        proveedor_id = request.form.get("proveedor_id") or None
        codigo = (producto.get("codigo_barras") or "").strip()
        if not codigo:
            codigo = generar_codigo_barra("PROD")

        foto_archivo = request.files.get("foto")
        reemplaza_foto = foto_archivo and foto_archivo.filename
        quitar_foto = request.form.get("quitar_foto") == "1" and not reemplaza_foto
        nueva_foto_drive = None

        try:
            if reemplaza_foto:
                nueva_foto_drive = subir_foto_drive(foto_archivo, "producto", id, nombre)
                reconectar_mysql_si_hace_falta()

            cursor.execute("""
                UPDATE productos
                SET nombre=%s,
                    precio=%s,
                    stock=%s,
                    proveedor_id=%s,
                    codigo_barras=%s
                WHERE id=%s
            """, (nombre, precio, stock, proveedor_id, codigo, id))

            if nueva_foto_drive:
                cursor.execute("""
                    UPDATE productos
                    SET foto_drive_id=%s,
                        foto_drive_nombre=%s,
                        foto_mime_type=%s,
                        foto_drive_web_view_url=%s,
                        foto_actualizada=NOW()
                    WHERE id=%s
                """, (
                    nueva_foto_drive["foto_drive_id"],
                    nueva_foto_drive["foto_drive_nombre"],
                    nueva_foto_drive["foto_mime_type"],
                    nueva_foto_drive["foto_drive_web_view_url"],
                    id
                ))
            elif quitar_foto:
                cursor.execute("""
                    UPDATE productos
                    SET foto_drive_id=NULL,
                        foto_drive_nombre=NULL,
                        foto_mime_type=NULL,
                        foto_drive_web_view_url=NULL,
                        foto_actualizada=NULL
                    WHERE id=%s
                """, (id,))

            mysql.connection.commit()
        except (ValueError, RuntimeError, MySQLdb.OperationalError) as error:
            mysql.connection.rollback()
            if nueva_foto_drive:
                eliminar_foto_drive_silencioso(nueva_foto_drive["foto_drive_id"])
            flash(str(error))
            cursor.close()
            return redirect(url_for("editar_producto", id=id))

        if nueva_foto_drive or quitar_foto:
            eliminar_foto_drive_silencioso(producto.get("foto_drive_id"))

        cursor.close()
        flash("Producto actualizado correctamente")
        return redirect(url_for("productos"))

    cursor.execute("SELECT * FROM proveedores ORDER BY nombre")
    proveedores = cursor.fetchall()
    cursor.close()
    return render_template("editar_producto.html", producto=producto, proveedores=proveedores)

# EXPORTAR PRODUCTO A EXCEL
# -------------------
@app.route("/exportar_productos")
@login_required
@rol_requerido(["admin"])
def exportar_productos():

    rellenar_codigos_barras_faltantes("productos")
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("""
        SELECT productos.*, proveedores.nombre AS proveedor_nombre
        FROM productos
        LEFT JOIN proveedores ON productos.proveedor_id = proveedores.id
    """)
    productos = cursor.fetchall()
    cursor.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"

    ws.append(["ID", "Nombre", "Stock", "Precio", "Proveedor", "Código Barras"])

    for p in productos:
        ws.append([
            p["id"],
            p["nombre"],
            p["stock"],
            float(p["precio"]),
            p.get("proveedor_nombre", ""),
            p.get("codigo_barras", "")
        ])

    file = io.BytesIO()
    wb.save(file)
    file.seek(0)

    return send_file(
        file,
        download_name="productos.xlsx",
        as_attachment=True
    )

# ETIQUETAS DE PRODUCTOS EN PDF
# -------------------
@app.route("/etiquetas_productos_carta")
@login_required
@rol_requerido(["admin"])
def etiquetas_productos_carta():

    rellenar_codigos_barras_faltantes("productos")
    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=10,
        leftMargin=10,
        topMargin=10,
        bottomMargin=10
    )

    elements = []
    styles = getSampleStyleSheet()

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("SELECT nombre, precio, codigo_barras FROM productos")
    productos = cursor.fetchall()
    cursor.close()

    etiquetas = []
    fila = []

    for p in productos:

        contenido = []

        contenido.append(Paragraph(f"<b>{p['nombre']}</b>", styles["Normal"]))
        contenido.append(Spacer(1, 3))
        contenido.append(Paragraph(f"${p['precio']}", styles["Normal"]))
        contenido.append(Spacer(1, 4))

        if p["codigo_barras"]:
            barcode = code128.Code128(
                str(p["codigo_barras"]),
                barHeight=15*mm,
                barWidth=1.0
            )
            contenido.append(barcode)
            contenido.append(Spacer(1, 3))
            contenido.append(Paragraph(str(p["codigo_barras"]), styles["Normal"]))

        etiqueta = Table([[contenido]],
                         colWidths=60*mm,
                         rowHeights=38*mm)

        etiqueta.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.black),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ]))

        fila.append(etiqueta)

        if len(fila) == 3:
            etiquetas.append(fila)
            fila = []

    if fila:
        while len(fila) < 3:
            fila.append("")
        etiquetas.append(fila)

    tabla_final = Table(etiquetas, colWidths=[65*mm]*3)
    elements.append(tabla_final)

    doc.build(elements)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="etiquetas_productos_carta.pdf",
        mimetype="application/pdf"
    )

# ETIQUETA DE UN SOLO PRODUCTO (IMPRESORA TERMICA)
# -------------------
@app.route("/etiqueta_producto/<int:id>")
@login_required
@rol_requerido(["admin"])
def etiqueta_producto(id):

    codigo = asegurar_codigo_barra_actual("productos", id)

    if not codigo:
        flash("Producto no encontrado")
        return redirect(url_for("productos"))

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("SELECT nombre, precio, codigo_barras FROM productos WHERE id=%s", (id,))
    producto = cursor.fetchone()
    cursor.close()

    return render_template("etiqueta_item_print.html",
                           nombre=producto["nombre"],
                           precio=producto["precio"],
                           codigo_barras=producto["codigo_barras"])


# ETIQUETA PNG DE UN SOLO PRODUCTO (IMPRESORA TERMICA BLUETOOTH)
# -------------------
@app.route("/etiqueta_producto/<int:id>/png")
@login_required
@rol_requerido(["admin"])
def etiqueta_producto_png(id):

    codigo = asegurar_codigo_barra_actual("productos", id)

    if not codigo:
        flash("Producto no encontrado")
        return redirect(url_for("productos"))

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("SELECT nombre, precio, codigo_barras FROM productos WHERE id=%s", (id,))
    producto = cursor.fetchone()
    cursor.close()

    imagen = generar_etiqueta_png(producto["nombre"], producto["precio"], producto["codigo_barras"])

    return send_file(
        imagen,
        # application/octet-stream (en vez de image/png): en WebViews de Android
        # viejos (ej. Lollipop 5.0) el motor intenta renderizar image/png inline
        # como si fuera la pagina en si, ignorando Content-Disposition:
        # attachment, en vez de entregarlo al DownloadListener nativo. Un tipo
        # generico fuerza a cualquier WebView a tratarlo como descarga.
        mimetype="application/octet-stream",
        as_attachment=True,
        download_name=f"etiqueta_producto_{id}.png"
    )


@app.route("/producto/<int:id>/foto")
@login_required
def foto_producto(id):
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("""
        SELECT foto_drive_id, foto_drive_nombre, foto_mime_type
        FROM productos
        WHERE id=%s
    """, (id,))
    producto = cursor.fetchone()
    cursor.close()

    if not producto or not producto.get("foto_drive_id"):
        abort(404)

    try:
        archivo = descargar_foto_drive(producto["foto_drive_id"])
    except RuntimeError:
        abort(404)

    return send_file(
        archivo,
        mimetype=producto.get("foto_mime_type") or "image/jpeg",
        download_name=producto.get("foto_drive_nombre") or f"producto-{id}.jpg",
        max_age=3600
    )

# -------------------
# CITAS
# -------------------

@app.route("/citas", methods=["GET", "POST"])
@login_required
def citas():

    actualizar_citas_perdidas()
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    if request.method == "POST":
        mascota_opcion = request.form.get("mascota_opcion", "existente")
        mascota_id = request.form.get("mascota_id")
        fecha = request.form["fecha"]
        servicio_id = request.form.get("servicio_id") or None
        notas = request.form.get("notas")
        fecha_convertida = datetime.strptime(fecha, "%Y-%m-%dT%H:%M")

        if mascota_opcion == "nueva":
            mascota_nombre = (request.form.get("mascota_nombre") or "").strip()
            cliente_id = request.form.get("cliente_id") or None

            if not mascota_nombre:
                flash("El nombre de la nueva mascota es obligatorio")
                cursor.close()
                return redirect(url_for("citas"))

            if not cliente_id:
                cliente_nombre = (request.form.get("cliente_nombre") or "").strip()
                cliente_telefono = (request.form.get("cliente_telefono") or "").strip()
                cliente_direccion = (request.form.get("cliente_direccion") or "").strip()

                if not cliente_nombre:
                    flash("Selecciona un dueño o escribe el nombre del dueño nuevo")
                    cursor.close()
                    return redirect(url_for("citas"))

                cursor.execute("""
                    INSERT INTO clientes (nombre, telefono, direccion)
                    VALUES (%s, %s, %s)
                """, (cliente_nombre, cliente_telefono, cliente_direccion))
                mysql.connection.commit()
                cliente_id = cursor.lastrowid

            cursor.execute("""
                INSERT INTO mascotas
                (nombre, especie, cliente_id, raza, edad, genero, talla, color, alergias, notas)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """, (
                mascota_nombre,
                request.form.get("mascota_especie") or "Perro",
                cliente_id,
                request.form.get("mascota_raza"),
                request.form.get("mascota_edad") or None,
                request.form.get("mascota_genero"),
                request.form.get("mascota_talla"),
                request.form.get("mascota_color"),
                request.form.get("mascota_alergias"),
                request.form.get("mascota_notas"),
            ))
            mysql.connection.commit()
            mascota_id = cursor.lastrowid
        else:
            if not mascota_id:
                flash("Selecciona una mascota de la lista")
                cursor.close()
                return redirect(url_for("citas"))

        cursor.execute("""
            INSERT INTO citas (mascota_id, servicio_id, fecha, estado, notas)
            VALUES (%s, %s, %s, 'Pendiente', %s)
        """, (mascota_id, servicio_id, fecha_convertida, notas))
        mysql.connection.commit()
        flash("Cita agendada correctamente")

    cursor.execute("""
        SELECT citas.*, mascotas.nombre as mascota, servicios.nombre as servicio
        FROM citas
        JOIN mascotas ON citas.mascota_id = mascotas.id
        LEFT JOIN servicios ON citas.servicio_id = servicios.id
        ORDER BY citas.fecha DESC
    """)
    data = cursor.fetchall()

    cursor.execute("""
        SELECT mascotas.*, clientes.nombre as dueno
        FROM mascotas
        LEFT JOIN clientes ON mascotas.cliente_id = clientes.id
        ORDER BY mascotas.nombre
    """)
    mascotas = cursor.fetchall()

    cursor.execute("SELECT * FROM servicios ORDER BY nombre")
    servicios = cursor.fetchall()

    cursor.execute("SELECT * FROM clientes ORDER BY nombre")
    clientes = cursor.fetchall()

    agenda_hoy = resumen_agenda_dia(cursor, datetime.now())

    cursor.close()

    return render_template("citas.html",
                           citas=data,
                           mascotas=mascotas,
                           servicios=servicios,
                           clientes=clientes,
                           agenda_hoy=agenda_hoy)


@app.route("/api/agenda_dia")
@login_required
def api_agenda_dia():
    fecha_raw = request.args.get("fecha")
    mascota_id = request.args.get("mascota_id") or None

    if not fecha_raw:
        return {"ok": False, "error": "Fecha requerida"}

    try:
        fecha = datetime.strptime(fecha_raw, "%Y-%m-%dT%H:%M")
    except ValueError:
        return {"ok": False, "error": "Formato de fecha inválido"}

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    agenda = resumen_agenda_dia(cursor, fecha, mascota_id)
    cursor.close()

    return {
        "ok": True,
        "ocupado_min": agenda["ocupado_min"],
        "disponible_min": agenda["disponible_min"],
        "nuevo_min": agenda["nuevo_min"],
        "cabe": agenda["cabe"],
    }

# EDITAR CITA
# -------------------
@app.route("/editar_cita/<int:id>", methods=["GET", "POST"])
@login_required
@rol_requerido(["admin", "empleado"])
def editar_cita(id):

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cursor.execute("SELECT * FROM citas WHERE id=%s", (id,))
    cita = cursor.fetchone()

    if not cita:
        cursor.close()
        flash("Cita no encontrada")
        return redirect(url_for("citas"))

    if request.method == "POST":

        fecha_raw = request.form["fecha"]
        servicio_id = request.form.get("servicio_id") or None
        notas = request.form.get("notas")
        observaciones = request.form.get("observaciones")
        tiempo_servicio_min = request.form.get("tiempo_servicio_min") or None

        # convertir formato HTML a formato MySQL
        fecha_convertida = datetime.strptime(
            fecha_raw, "%Y-%m-%dT%H:%M"
        )

        estado = request.form["estado"]

        cursor.execute("""
            UPDATE citas
            SET fecha=%s,
                servicio_id=%s,
                estado=%s,
                notas=%s,
                observaciones=%s,
                tiempo_servicio_min=%s
            WHERE id=%s
        """, (fecha_convertida, servicio_id, estado, notas, observaciones, tiempo_servicio_min, id))

        mysql.connection.commit()
        flash("Cita actualizada correctamente")
        return redirect(url_for("citas"))

    cursor.close()

    # convertir fecha MySQL a formato HTML datetime-local
    fecha_html = cita["fecha"].strftime("%Y-%m-%dT%H:%M")

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("SELECT * FROM servicios ORDER BY nombre")
    servicios = cursor.fetchall()
    cursor.close()

    return render_template("editar_cita.html",
                           cita=cita,
                           fecha_html=fecha_html,
                           servicios=servicios)

# CANCELAR CITA
# -------------------
@app.route("/cancelar_cita/<int:id>")
@login_required
def cancelar_cita(id):

    cursor = mysql.connection.cursor()

    cursor.execute("""
        UPDATE citas
        SET estado='Cancelada'
        WHERE id=%s
    """, (id,))

    mysql.connection.commit()
    cursor.close()

    flash("Cita cancelada correctamente")
    return redirect(url_for("citas"))

# CONVERTIR CITA A VENTA
# -------------------
@app.route("/convertir_cita/<int:id>", methods=["GET", "POST"])
@login_required
def convertir_cita(id):

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cursor.execute("""
        SELECT citas.*,
               mascotas.nombre as mascota,
               mascotas.cliente_id,
               clientes.nombre as cliente,
               servicios.nombre as servicio,
               servicios.precio
        FROM citas
        JOIN mascotas ON citas.mascota_id = mascotas.id
        LEFT JOIN clientes ON mascotas.cliente_id = clientes.id
        LEFT JOIN servicios ON citas.servicio_id = servicios.id
        WHERE citas.id=%s
    """, (id,))
    cita = cursor.fetchone()

    if not cita:
        cursor.close()
        flash("Cita no encontrada")
        return redirect(url_for("citas"))

    if request.method == "POST":
        observaciones = request.form.get("observaciones")
        tiempo = request.form.get("tiempo_servicio_min") or None
        citas_mismo_dueno_ids = []

        for cita_relacionada_id in request.form.getlist("citas_mismo_dueno"):
            try:
                citas_mismo_dueno_ids.append(int(cita_relacionada_id))
            except (TypeError, ValueError):
                continue

        cursor.execute("""
            UPDATE citas
            SET estado='Completada',
                observaciones=%s,
                tiempo_servicio_min=%s
            WHERE id=%s
        """, (observaciones, tiempo, id))

        citas_mismo_dueno = []
        if citas_mismo_dueno_ids and cita.get("cliente_id"):
            placeholders = ", ".join(["%s"] * len(citas_mismo_dueno_ids))
            params = tuple(citas_mismo_dueno_ids) + (id, cita["cliente_id"], cita["fecha"])

            cursor.execute(f"""
                SELECT c.id,
                       c.mascota_id,
                       c.servicio_id,
                       m.nombre AS mascota,
                       s.nombre AS servicio,
                       s.precio
                FROM citas c
                JOIN mascotas m ON c.mascota_id = m.id
                LEFT JOIN servicios s ON c.servicio_id = s.id
                WHERE c.id IN ({placeholders})
                AND c.id<>%s
                AND m.cliente_id=%s
                AND DATE(c.fecha)=DATE(%s)
                AND c.estado='Pendiente'
            """, params)
            citas_mismo_dueno = cursor.fetchall()

            # Cada mascota conserva su propio tiempo/observaciones: no se
            # reutilizan los de la cita principal para no mezclar datos.
            for cita_relacionada in citas_mismo_dueno:
                cid = cita_relacionada["id"]
                obs_relacionada = request.form.get(f"observaciones_{cid}")
                tiempo_relacionado = request.form.get(f"tiempo_servicio_min_{cid}") or None
                cursor.execute("""
                    UPDATE citas
                    SET estado='Completada',
                        observaciones=%s,
                        tiempo_servicio_min=%s
                    WHERE id=%s
                """, (obs_relacionada, tiempo_relacionado, cid))

        mysql.connection.commit()

        session["venta_mascota_id"] = cita["mascota_id"]
        session["venta_cita_id"] = cita["id"]

        if "carrito" not in session:
            session["carrito"] = []

        if cita.get("servicio_id"):
            session["carrito"].append({
                "tipo": "servicio",
                "id": cita["servicio_id"],
                "nombre": cita["servicio"],
                "precio": float(cita["precio"] or 0),
                "cantidad": 1,
                "subtotal": float(cita["precio"] or 0),
                "mascota_id": cita["mascota_id"],
                "mascota_nombre": cita["mascota"],
            })

        for cita_relacionada in citas_mismo_dueno:
            if cita_relacionada.get("servicio_id"):
                session["carrito"].append({
                    "tipo": "servicio",
                    "id": cita_relacionada["servicio_id"],
                    "nombre": cita_relacionada["servicio"],
                    "precio": float(cita_relacionada["precio"] or 0),
                    "cantidad": 1,
                    "subtotal": float(cita_relacionada["precio"] or 0),
                    "mascota_id": cita_relacionada["mascota_id"],
                    "mascota_nombre": cita_relacionada["mascota"],
                })

        session.modified = True

        flash("Cita completada. Revisa el carrito y finaliza el cobro.")
        cursor.close()
        return redirect(url_for("ventas"))

    citas_mismo_dueno = []
    if cita.get("cliente_id"):
        cursor.execute("""
            SELECT c.id,
                   c.fecha,
                   m.nombre AS mascota,
                   s.nombre AS servicio
            FROM citas c
            JOIN mascotas m ON c.mascota_id = m.id
            LEFT JOIN servicios s ON c.servicio_id = s.id
            WHERE c.id<>%s
            AND m.cliente_id=%s
            AND DATE(c.fecha)=DATE(%s)
            AND c.estado='Pendiente'
            ORDER BY c.fecha
        """, (id, cita["cliente_id"], cita["fecha"]))
        citas_mismo_dueno = cursor.fetchall()

    cursor.close()
    return render_template("finalizar_cita.html",
                           cita=cita,
                           citas_mismo_dueno=citas_mismo_dueno)

# -------------------
# VENTAS
# -------------------
@app.route("/ventas", methods=["GET", "POST"])
@login_required
def ventas():

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    def cerrar_y_redirigir(endpoint, **kwargs):
        cursor.close()
        return redirect(url_for(endpoint, **kwargs))

    if "carrito" not in session:
        session["carrito"] = []

    # ==========================
    # AGREGAR AL CARRITO
    # ==========================
    if request.method == "POST" and "agregar" in request.form:

        codigo = request.form.get("codigo_barras")
        cantidad = int(request.form.get("cantidad", 1))

        if not codigo:
            flash("Código no proporcionado")
            return cerrar_y_redirigir("ventas")

        # Buscar en productos
        cursor.execute("SELECT * FROM productos WHERE codigo_barras=%s", (codigo,))
        item = cursor.fetchone()

        if item:
            tipo = "producto"

            if item["stock"] < cantidad:
                flash("Stock insuficiente")
                return cerrar_y_redirigir("ventas")

        else:
            # Buscar en servicios
            cursor.execute("SELECT * FROM servicios WHERE codigo_barras=%s", (codigo,))
            item = cursor.fetchone()

            if item:
                tipo = "servicio"
            else:
                flash("Código no encontrado")
                return cerrar_y_redirigir("ventas")

        carrito_item = {
            "tipo": tipo,
            "id": item["id"],
            "nombre": item["nombre"],
            "precio": float(item["precio"]),
            "cantidad": cantidad,
            "subtotal": float(item["precio"]) * cantidad
        }

        session["carrito"].append(carrito_item)
        session.modified = True

        return cerrar_y_redirigir("ventas")

    # ==========================
    # ELIMINAR ITEM
    # ==========================
    if request.method == "POST" and "eliminar" in request.form:

        index = int(request.form.get("index"))
        session["carrito"].pop(index)
        session.modified = True

        return cerrar_y_redirigir("ventas")

    # ==========================
    # VACIAR CARRITO
    # ==========================
    if request.method == "POST" and "vaciar" in request.form:

        session["carrito"] = []
        session.modified = True

        return cerrar_y_redirigir("ventas")

    # ==========================
    # FINALIZAR VENTA
    # ==========================
    if request.method == "POST" and "finalizar" in request.form:

        carrito = session.get("carrito", [])

        if not carrito:
            flash("Carrito vacío")
            return cerrar_y_redirigir("ventas")

        metodo_pago = request.form.get("metodo_pago")
        cliente_id = request.form.get("cliente_id")
        mascota_id = request.form.get("mascota_id") or session.get("venta_mascota_id")
        mascota_opcion = request.form.get("mascota_opcion")
        cita_id = session.get("venta_cita_id")

        if mascota_opcion == "nueva":
            mascota_nombre = request.form.get("mascota_nombre")

            if not mascota_nombre:
                flash("El nombre de la mascota nueva es obligatorio")
                return cerrar_y_redirigir("ventas")

            cursor.execute("""
                INSERT INTO mascotas
                (nombre, especie, cliente_id, raza, edad, genero, talla, color, alergias, notas)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """, (
                mascota_nombre,
                request.form.get("mascota_especie") or "Perro",
                cliente_id if cliente_id else None,
                request.form.get("mascota_raza"),
                request.form.get("mascota_edad") or None,
                request.form.get("mascota_genero"),
                request.form.get("mascota_talla"),
                request.form.get("mascota_color"),
                request.form.get("mascota_alergias"),
                request.form.get("mascota_notas"),
            ))
            mysql.connection.commit()
            mascota_id = cursor.lastrowid

        if mascota_id and not cliente_id:
            cursor.execute("SELECT cliente_id FROM mascotas WHERE id=%s", (mascota_id,))
            mascota_venta = cursor.fetchone()
            if mascota_venta and mascota_venta.get("cliente_id"):
                cliente_id = mascota_venta["cliente_id"]

        if cliente_id and mascota_id and mascota_opcion != "nueva":
            cursor.execute(
                "SELECT id FROM mascotas WHERE id=%s AND cliente_id=%s",
                (mascota_id, cliente_id)
            )
            if not cursor.fetchone():
                flash("La mascota seleccionada no pertenece al cliente elegido")
                return cerrar_y_redirigir("ventas")

        total = sum(item["subtotal"] for item in carrito)

        cursor.execute(
            "INSERT INTO ventas (total, metodo_pago, cliente_id, mascota_id, cita_id) VALUES (%s, %s, %s, %s, %s)",
            (total, metodo_pago, cliente_id if cliente_id else None, mascota_id if mascota_id else None, cita_id)
        )

        mysql.connection.commit()
        venta_id = cursor.lastrowid

        for item in carrito:

            item_mascota_id = item.get("mascota_id") or mascota_id or None

            if item["tipo"] == "producto":

                cursor.execute("""
                    INSERT INTO detalle_venta
                    (venta_id, tipo, producto_id, servicio_id, cantidad, precio_unitario, subtotal, mascota_id)
                    VALUES (%s, %s, %s, NULL, %s, %s, %s, %s)
                """, (
                    venta_id,
                    "producto",
                    item["id"],
                    item["cantidad"],
                    item["precio"],
                    item["subtotal"],
                    item_mascota_id
                ))

                cursor.execute("""
                    UPDATE productos
                    SET stock = stock - %s
                    WHERE id = %s
                """, (item["cantidad"], item["id"]))

            elif item["tipo"] == "servicio":

                cursor.execute("""
                    INSERT INTO detalle_venta
                    (venta_id, tipo, producto_id, servicio_id, cantidad, precio_unitario, subtotal, mascota_id)
                    VALUES (%s, %s, NULL, %s, %s, %s, %s, %s)
                """, (
                    venta_id,
                    "servicio",
                    item["id"],
                    item["cantidad"],
                    item["precio"],
                    item["subtotal"],
                    item_mascota_id
                ))

        mysql.connection.commit()

        session["carrito"] = []
        session.pop("venta_mascota_id", None)
        session.pop("venta_cita_id", None)
        flash("Venta completada correctamente")

        return cerrar_y_redirigir("ticket_html", venta_id=venta_id)

    # ==========================
    # CARGAR DATOS
    # ==========================
    rellenar_codigos_barras_faltantes("productos")
    rellenar_codigos_barras_faltantes("servicios")

    cursor.execute("SELECT * FROM productos")
    productos = cursor.fetchall()

    cursor.execute("SELECT * FROM servicios")
    servicios = cursor.fetchall()

    cursor.execute("SELECT * FROM clientes")
    clientes = cursor.fetchall()

    cursor.execute("""
        SELECT mascotas.*, clientes.nombre as dueno
        FROM mascotas
        JOIN clientes ON mascotas.cliente_id = clientes.id
        ORDER BY mascotas.nombre
    """)
    mascotas = cursor.fetchall()

    venta_cliente_id = None
    if session.get("venta_mascota_id"):
        for mascota in mascotas:
            if mascota["id"] == session.get("venta_mascota_id"):
                venta_cliente_id = mascota.get("cliente_id")
                break

    cursor.close()

    return render_template(
        "ventas.html",
        productos=productos,
        servicios=servicios,
        clientes=clientes,
        mascotas=mascotas,
        venta_cliente_id=venta_cliente_id,
        carrito=session.get("carrito", [])
    )

# -------------------
# Cancelar Venta y Devolver Stock
# -------------------
@app.route("/cancelar_venta/<int:venta_id>")
@login_required
@rol_requerido(["admin"])
def cancelar_venta(venta_id):

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    # verificar venta
    cursor.execute("""
        SELECT v.*, c.nombre as cliente, m.nombre as mascota
        FROM ventas v
        LEFT JOIN clientes c ON v.cliente_id = c.id
        LEFT JOIN mascotas m ON v.mascota_id = m.id
        WHERE v.id=%s
    """, (venta_id,))
    venta = cursor.fetchone()

    if not venta:
        flash("Venta no encontrada")
        return redirect(url_for("ventas_lista"))

    if venta["estado"] == "cancelada":
        flash("La venta ya está cancelada")
        return redirect(url_for("ventas_lista"))

    # obtener detalle
    cursor.execute("""
        SELECT *
        FROM detalle_venta
        WHERE venta_id=%s
    """, (venta_id,))

    detalles = cursor.fetchall()

    # devolver inventario solo productos
    for d in detalles:

        if d["tipo"] == "producto":

            cursor.execute("""
                UPDATE productos
                SET stock = stock + %s
                WHERE id = %s
            """, (d["cantidad"], d["producto_id"]))

    # actualizar estado venta
    cursor.execute("""
        UPDATE ventas
        SET estado='cancelada'
        WHERE id=%s
    """, (venta_id,))

    mysql.connection.commit()
    cursor.close()

    flash("Venta cancelada y stock devuelto correctamente")

    return redirect(url_for("ventas_lista"))

# -------------------
# Ticket de venta
# -------------------
@app.route("/ticket/<int:venta_id>")
@login_required
def generar_ticket(venta_id):

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer)
    elements = []
    styles = getSampleStyleSheet()

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cursor.execute("""
        SELECT v.*, c.nombre as cliente, m.nombre as mascota
        FROM ventas v
        LEFT JOIN clientes c ON v.cliente_id = c.id
        LEFT JOIN mascotas m ON v.mascota_id = m.id
        WHERE v.id=%s
    """, (venta_id,))
    venta = cursor.fetchone()

    cursor.execute("""
    SELECT 
        d.tipo,
        d.cantidad,
        d.subtotal,
        COALESCE(p.nombre, s.nombre) as nombre
    FROM detalle_venta d
    LEFT JOIN productos p ON d.producto_id = p.id
    LEFT JOIN servicios s ON d.servicio_id = s.id
    WHERE d.venta_id = %s
""", (venta_id,))

    detalles = cursor.fetchall()
    cursor.close()

    elements.append(Paragraph(f"Ticket Venta #{venta_id}", styles["Title"]))
    elements.append(Spacer(1, 12))

    for d in detalles:
        elements.append(Paragraph(
            f"{d['nombre']} x{d['cantidad']} - ${d['subtotal']}",
            styles["Normal"]
        ))

    elements.append(Spacer(1, 12))
    elements.append(Paragraph(f"Total: ${venta['total']}", styles["Heading2"]))

    doc.build(elements)
    buffer.seek(0)

    return send_file(buffer,
                     as_attachment=True,
                     download_name=f"ticket_{venta_id}.pdf",
                     mimetype="application/pdf")

# -------------------
# Impresión de Ticket 58mm / 80mm
# -------------------
def cargar_ticket_venta(venta_id):
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cursor.execute("""
        SELECT v.*,
               c.nombre AS cliente,
               m.nombre AS mascota
        FROM ventas v
        LEFT JOIN clientes c ON v.cliente_id = c.id
        LEFT JOIN mascotas m ON v.mascota_id = m.id
        WHERE v.id=%s
    """, (venta_id,))
    venta = cursor.fetchone()

    cursor.execute("""
        SELECT
            d.tipo,
            d.cantidad,
            d.subtotal,
            COALESCE(p.nombre, s.nombre) as nombre
        FROM detalle_venta d
        LEFT JOIN productos p ON d.producto_id=p.id
        LEFT JOIN servicios s ON d.servicio_id=s.id
        WHERE venta_id=%s
    """, (venta_id,))
    detalles = cursor.fetchall()
    cursor.close()

    return venta, detalles


@app.route("/ticket_html/<int:venta_id>")
@login_required
def ticket_html(venta_id):
    return redirect(url_for("ticket_html_copia", venta_id=venta_id, copia="cliente"))


@app.route("/ticket_html/<int:venta_id>/<copia>")
@login_required
def ticket_html_copia(venta_id, copia):
    if copia not in ("cliente", "negocio"):
        abort(404)

    venta, detalles = cargar_ticket_venta(venta_id)
    if not venta:
        flash("Venta no encontrada")
        return redirect(url_for("ventas"))

    if copia == "cliente":
        siguiente_url = url_for("ticket_html_copia", venta_id=venta_id, copia="negocio")
        return render_template("ticket_print_cliente.html",
                               venta=venta,
                               detalles=detalles,
                               siguiente_url=siguiente_url)

    siguiente_url = url_for("ventas")
    if request.args.get("siguiente") == "reporte":
        fecha_reporte = request.args.get("fecha")
        if fecha_reporte:
            siguiente_url = url_for("reporte_diario", fecha=fecha_reporte)
        else:
            siguiente_url = url_for("reporte_diario")

    return render_template("ticket_print_negocio.html",
                           venta=venta,
                           detalles=detalles,
                           siguiente_url=siguiente_url)

# -------------------
# REPORTE DE VENTAS DIARIO
# -------------------
@app.route("/reporte_diario")
@login_required
@rol_requerido(["admin"])
def reporte_diario():

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    fecha_raw = request.args.get("fecha") or datetime.now().strftime("%Y-%m-%d")

    try:
        fecha_reporte = datetime.strptime(fecha_raw, "%Y-%m-%d").date()
    except ValueError:
        fecha_reporte = datetime.now().date()
        fecha_raw = fecha_reporte.strftime("%Y-%m-%d")

    cursor.execute("""
        SELECT
            IFNULL(SUM(total),0) as total_dia,
            COUNT(*) as ventas
        FROM ventas
        WHERE DATE(fecha)=%s
        AND estado='activa'
    """, (fecha_reporte,))

    data = cursor.fetchone()

    cursor.execute("""
        SELECT metodo_pago,
               IFNULL(SUM(total),0) as total,
               COUNT(*) as ventas
        FROM ventas
        WHERE DATE(fecha)=%s
        AND estado='activa'
        GROUP BY metodo_pago
        ORDER BY metodo_pago
    """, (fecha_reporte,))
    totales_pago = cursor.fetchall()

    cursor.execute("""
        SELECT v.id,
               v.fecha,
               v.total,
               v.metodo_pago,
               c.nombre AS cliente,
               m.nombre AS mascota
        FROM ventas v
        LEFT JOIN clientes c ON v.cliente_id = c.id
        LEFT JOIN mascotas m ON v.mascota_id = m.id
        WHERE DATE(v.fecha)=%s
        AND v.estado='activa'
        ORDER BY v.fecha DESC
    """, (fecha_reporte,))
    tickets = cursor.fetchall()

    cursor.close()

    return render_template("reporte.html",
                           data=data,
                           totales_pago=totales_pago,
                           tickets=tickets,
                           fecha_reporte=fecha_raw)


@app.route("/reporte_diario/descargar")
@login_required
@rol_requerido(["admin"])
def descargar_reporte_diario():
    fecha_raw = request.args.get("fecha") or datetime.now().strftime("%Y-%m-%d")

    try:
        fecha_reporte = datetime.strptime(fecha_raw, "%Y-%m-%d").date()
    except ValueError:
        fecha_reporte = datetime.now().date()
        fecha_raw = fecha_reporte.strftime("%Y-%m-%d")

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("""
        SELECT v.id,
               v.fecha,
               v.total,
               v.metodo_pago,
               c.nombre AS cliente,
               m.nombre AS mascota
        FROM ventas v
        LEFT JOIN clientes c ON v.cliente_id = c.id
        LEFT JOIN mascotas m ON v.mascota_id = m.id
        WHERE DATE(v.fecha)=%s
        AND v.estado='activa'
        ORDER BY v.fecha ASC
    """, (fecha_reporte,))
    tickets = cursor.fetchall()

    cursor.execute("""
        SELECT metodo_pago,
               IFNULL(SUM(total),0) as total,
               COUNT(*) as ventas
        FROM ventas
        WHERE DATE(fecha)=%s
        AND estado='activa'
        GROUP BY metodo_pago
        ORDER BY metodo_pago
    """, (fecha_reporte,))
    totales_pago = cursor.fetchall()
    cursor.close()

    salida = io.StringIO()
    writer = csv.writer(salida)
    writer.writerow(["Reporte diario", fecha_raw])
    writer.writerow([])
    writer.writerow(["Ticket", "Fecha", "Cliente", "Mascota", "Metodo de pago", "Total"])

    for ticket in tickets:
        writer.writerow([
            ticket["id"],
            ticket["fecha"],
            ticket.get("cliente") or "",
            ticket.get("mascota") or "",
            ticket.get("metodo_pago") or "",
            f"{float(ticket['total'] or 0):.2f}",
        ])

    writer.writerow([])
    writer.writerow(["Metodo de pago", "Ventas", "Total"])

    for total in totales_pago:
        writer.writerow([
            total.get("metodo_pago") or "Sin especificar",
            total["ventas"],
            f"{float(total['total'] or 0):.2f}",
        ])

    memoria = io.BytesIO()
    memoria.write(salida.getvalue().encode("utf-8-sig"))
    memoria.seek(0)

    return send_file(memoria,
                     as_attachment=True,
                     download_name=f"reporte_diario_{fecha_raw}.csv",
                     mimetype="text/csv")


# -------------------
# CORTE DE CAJA
# -------------------
@app.route("/corte")
@login_required
def corte():

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cursor.execute("""
        SELECT SUM(total) as total,
               COUNT(*) as ventas
        FROM ventas
        WHERE DATE(fecha) = CURDATE()
    """)

    datos = cursor.fetchone()

    cursor.execute("""
        INSERT INTO corte_caja (total_ventas, cantidad_ventas)
        VALUES (%s, %s)
    """, (datos["total"] or 0, datos["ventas"] or 0))

    mysql.connection.commit()
    cursor.close()

    flash("Corte realizado correctamente")
    return redirect(url_for("dashboard"))


# Abrir caja con monto inicial 
# -------------------
@app.route("/abrir_caja", methods=["GET", "POST"])
@login_required
def abrir_caja():

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cerrar_cajas_vencidas(cursor)
    mysql.connection.commit()

    cursor.execute("""
        SELECT *
        FROM caja
        WHERE estado='abierta'
        AND DATE(fecha_apertura)=CURDATE()
        ORDER BY fecha_apertura DESC
        LIMIT 1
    """)
    caja_abierta = cursor.fetchone()

    if caja_abierta:
        cursor.close()
        if session.get("rol") == "admin":
            flash("Ya existe una caja abierta")
        return redirect(url_for("dashboard"))

    if request.method == "POST":

        monto = float(request.form.get("monto_inicial"))

        cursor.execute("""
            INSERT INTO caja (monto_inicial)
            VALUES (%s)
        """, (monto,))

        mysql.connection.commit()
        flash("Caja abierta correctamente")
        cursor.close()
        return redirect(url_for("dashboard"))

    cursor.close()
    return render_template("abrir_caja.html",
                           caja_requerida=request.args.get("requerida") == "1")


# Cerrar caja con monto final y diferencias
# -------------------
@app.route("/cerrar_caja", methods=["GET", "POST"])
@login_required
def cerrar_caja():

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    # Buscar caja abierta
    cursor.execute("""
        SELECT *
        FROM caja
        WHERE estado='abierta'
        AND DATE(fecha_apertura)=CURDATE()
        ORDER BY fecha_apertura DESC
        LIMIT 1
    """)
    caja = cursor.fetchone()

    if not caja:
        cursor.close()
        flash("No hay caja abierta")
        return redirect(url_for("dashboard"))

    if request.method == "POST":

        try:
            monto_final = float(request.form.get("monto_final"))

            # Calcular ventas desde apertura hasta ahora
            cursor.execute("""
                SELECT IFNULL(SUM(total),0) as total_ventas
                FROM ventas
                WHERE estado='activa'
                AND fecha BETWEEN %s AND NOW()
            """, (caja["fecha_apertura"],))

            resultado = cursor.fetchone()
            total_ventas = resultado["total_ventas"] if resultado else 0

            diferencia = monto_final - (float(caja["monto_inicial"]) + float(total_ventas))

            cursor.execute("""
                UPDATE caja
                SET fecha_cierre=NOW(),
                    monto_final=%s,
                    diferencia=%s,
                    estado='cerrada'
                WHERE id=%s
            """, (monto_final, diferencia, caja["id"]))

            mysql.connection.commit()
            session.pop("caja_verificada", None)

            if session.get("rol") == "admin":
                flash(f"Caja cerrada correctamente. Diferencia: ${diferencia}")
            else:
                flash("Caja cerrada correctamente")

        except Exception as e:
            mysql.connection.rollback()
            flash(f"Error al cerrar caja: {str(e)}")

        cursor.close()
        return redirect(url_for("dashboard"))

    cursor.close()
    return render_template("cerrar_caja.html", caja=caja)

# -------------------
# LISTA DE VENTAS
# -------------------
@app.route("/ventas_lista")
@login_required
def ventas_lista():

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cursor.execute("""
    SELECT id, fecha, total, metodo_pago, estado
    FROM ventas
    ORDER BY fecha DESC
""")

    ventas = cursor.fetchall()
    cursor.close()

    return render_template("ventas_lista.html", ventas=ventas)

# -------------------
# DETALLES DE CADA VENTA
@app.route("/venta/<int:venta_id>")
@login_required
def ver_venta(venta_id):

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cursor.execute("SELECT * FROM ventas WHERE id=%s", (venta_id,))
    venta = cursor.fetchone()

    cursor.execute("""
        SELECT 
            d.tipo,
            d.cantidad,
            d.subtotal,
            COALESCE(p.nombre, s.nombre) as nombre
        FROM detalle_venta d
        LEFT JOIN productos p ON d.producto_id = p.id
        LEFT JOIN servicios s ON d.servicio_id = s.id
        WHERE d.venta_id = %s
    """, (venta_id,))

    detalles = cursor.fetchall()

    cursor.close()

    return render_template("ver_venta.html",
                           venta=venta,
                           detalles=detalles)


# -------------------
# Entrada de Inventario con lector de codigo de barras 
# -------------------
@app.route("/entrada_inventario", methods=["GET", "POST"])
@login_required
def entrada_inventario():

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    if request.method == "POST":

        codigo = request.form.get("codigo_barras")
        cantidad = int(request.form.get("cantidad", 1))

        cursor.execute("SELECT * FROM productos WHERE codigo_barras=%s", (codigo,))
        producto = cursor.fetchone()

        if not producto:
            flash("Producto no encontrado")
            return redirect(url_for("entrada_inventario"))

        # actualizar stock
        cursor.execute("""
            UPDATE productos
            SET stock = stock + %s
            WHERE id = %s
        """, (cantidad, producto["id"]))

        # registrar entrada
        cursor.execute("""
            INSERT INTO entradas_inventario (producto_id, cantidad)
            VALUES (%s, %s)
        """, (producto["id"], cantidad))

        mysql.connection.commit()

        flash("Entrada registrada correctamente")

        return redirect(url_for("entrada_inventario"))

    cursor.close()

    return render_template("entrada_inventario.html")


# -------------------
# SERVICIOS
# -------------------
@app.route("/servicios", methods=["GET", "POST"])
@login_required
def servicios():

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    if request.method == "POST":

        nombre = request.form.get("nombre")
        precio = request.form.get("precio")
        descripcion = request.form.get("descripcion")
        codigo = generar_codigo_barra("SERV")

        cursor.execute("""
            INSERT INTO servicios (nombre, precio, descripcion, codigo_barras)
            VALUES (%s, %s, %s, %s)
        """, (nombre, precio, descripcion, codigo))

        mysql.connection.commit()
        flash("Servicio agregado correctamente")

        return redirect(url_for("servicios"))

    rellenar_codigos_barras_faltantes("servicios")
    cursor.execute("SELECT * FROM servicios")
    servicios = cursor.fetchall()

    cursor.close()

    return render_template("servicios.html", servicios=servicios)

# EDITAR SERVICIO
# -------------------
@app.route("/editar_servicio/<int:id>", methods=["GET", "POST"])
@login_required
@rol_requerido(["admin"])
def editar_servicio(id):

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cursor.execute("SELECT * FROM servicios WHERE id=%s", (id,))
    servicio = cursor.fetchone()

    if not servicio:
        flash("Servicio no encontrado")
        return redirect(url_for("servicios"))

    if request.method == "POST":

        nombre = request.form["nombre"]
        precio = request.form["precio"]
        descripcion = request.form["descripcion"]
        codigo = (servicio.get("codigo_barras") or "").strip()
        if not codigo:
            codigo = generar_codigo_barra("SERV")

        cursor.execute("""
            UPDATE servicios
            SET nombre=%s,
                precio=%s,
                descripcion=%s,
                codigo_barras=%s
            WHERE id=%s
        """, (nombre, precio, descripcion, codigo, id))

        mysql.connection.commit()
        flash("Servicio actualizado correctamente")
        return redirect(url_for("servicios"))

    cursor.close()
    return render_template("editar_servicio.html", servicio=servicio)


# ELIMINAR SERVICIO
# -------------------
@app.route("/eliminar_servicio/<int:id>")
@login_required
@rol_requerido(["admin"])
def eliminar_servicio(id):

    cursor = mysql.connection.cursor()
    cursor.execute("DELETE FROM servicios WHERE id=%s", (id,))
    mysql.connection.commit()
    cursor.close()

    flash("Servicio eliminado correctamente")
    return redirect(url_for("servicios"))


# EXPORTAR SERVICIO A EXCEL
# -------------------
@app.route("/exportar_servicios")
@login_required
@rol_requerido(["admin"])
def exportar_servicios():

    rellenar_codigos_barras_faltantes("servicios")
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("SELECT * FROM servicios")
    servicios = cursor.fetchall()
    cursor.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Servicios"

    ws.append(["ID", "Nombre", "Precio", "Descripción", "Código Barras"])

    for s in servicios:
        ws.append([
            s["id"],
            s["nombre"],
            float(s["precio"]),
            s["descripcion"],
            s.get("codigo_barras", "")
        ])

    file = io.BytesIO()
    wb.save(file)
    file.seek(0)

    return send_file(
        file,
        download_name="servicios.xlsx",
        as_attachment=True
    )

# ETIQUETAS DE SERVICIOS EN PDF
# -------------------
@app.route("/etiquetas_servicios_carta")
@login_required
@rol_requerido(["admin"])
def etiquetas_servicios_carta():

    rellenar_codigos_barras_faltantes("servicios")
    buffer = io.BytesIO()

    doc = SimpleDocTemplate(buffer, pagesize=letter)

    elements = []
    styles = getSampleStyleSheet()

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("SELECT nombre, precio, codigo_barras FROM servicios")
    servicios = cursor.fetchall()
    cursor.close()

    etiquetas = []
    fila = []

    for s in servicios:

        contenido = []

        contenido.append(Paragraph(f"<b>{s['nombre']}</b>", styles["Normal"]))
        contenido.append(Spacer(1, 3))
        contenido.append(Paragraph(f"${s['precio']}", styles["Normal"]))
        contenido.append(Spacer(1, 4))

        if s["codigo_barras"]:
            barcode = code128.Code128(
                str(s["codigo_barras"]),
                barHeight=15*mm,
                barWidth=1.0
            )
            contenido.append(barcode)
            contenido.append(Spacer(1, 3))
            contenido.append(Paragraph(str(s["codigo_barras"]), styles["Normal"]))

        etiqueta = Table([[contenido]],
                         colWidths=60*mm,
                         rowHeights=38*mm)

        etiqueta.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.black),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ]))

        fila.append(etiqueta)

        if len(fila) == 3:
            etiquetas.append(fila)
            fila = []

    if fila:
        while len(fila) < 3:
            fila.append("")
        etiquetas.append(fila)

    elements.append(Table(etiquetas, colWidths=[65*mm]*3))

    doc.build(elements)
    buffer.seek(0)

    return send_file(buffer,
                     as_attachment=True,
                     download_name="etiquetas_servicios_carta.pdf",
                     mimetype="application/pdf")


# -------------------
# INSUMOS
# -------------------
@app.route("/insumos", methods=["GET", "POST"])
@login_required
@rol_requerido(["admin"])
def insumos():

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    if request.method == "POST":

        nombre = request.form["nombre"]
        stock = request.form["stock"]
        costo = request.form["costo"]
        proveedor_id = request.form.get("proveedor_id") or None

        try:
            # La foto se sube antes de escribir en la BD para no dejar un
            # insumo a medio insertar si Drive tarda y la conexion cae.
            foto_drive = subir_foto_drive(request.files.get("foto"), "insumo", "nuevo", nombre)
            reconectar_mysql_si_hace_falta()

            codigo = generar_codigo_barra("INS")

            cursor.execute("""
                INSERT INTO insumos (nombre, stock, costo, proveedor_id, codigo_barras)
                VALUES (%s, %s, %s, %s, %s)
            """, (nombre, stock, costo, proveedor_id, codigo))
            insumo_id = cursor.lastrowid

            if foto_drive:
                cursor.execute("""
                    UPDATE insumos
                    SET foto_drive_id=%s,
                        foto_drive_nombre=%s,
                        foto_mime_type=%s,
                        foto_drive_web_view_url=%s,
                        foto_actualizada=NOW()
                    WHERE id=%s
                """, (
                    foto_drive["foto_drive_id"],
                    foto_drive["foto_drive_nombre"],
                    foto_drive["foto_mime_type"],
                    foto_drive["foto_drive_web_view_url"],
                    insumo_id
                ))

            mysql.connection.commit()
            flash("Insumo agregado")
        except (ValueError, RuntimeError, MySQLdb.OperationalError) as error:
            mysql.connection.rollback()
            flash(str(error))

        cursor.close()
        return redirect(url_for("insumos"))

    rellenar_codigos_barras_faltantes("insumos")
    cursor.execute("""
        SELECT insumos.*, proveedores.nombre AS proveedor_nombre
        FROM insumos
        LEFT JOIN proveedores ON insumos.proveedor_id = proveedores.id
        ORDER BY insumos.fecha DESC
    """)
    insumos = cursor.fetchall()
    cursor.execute("SELECT * FROM proveedores ORDER BY nombre")
    proveedores = cursor.fetchall()
    cursor.close()

    return render_template("insumos.html", insumos=insumos, proveedores=proveedores)


# EDITAR INSUMO
# -------------------
@app.route("/editar_insumo/<int:id>", methods=["GET", "POST"])
@login_required
@rol_requerido(["admin"])
def editar_insumo(id):

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cursor.execute("SELECT * FROM insumos WHERE id=%s", (id,))
    insumo = cursor.fetchone()

    if not insumo:
        flash("Insumo no encontrado")
        return redirect(url_for("insumos"))

    if request.method == "POST":

        nombre = request.form["nombre"]
        stock = request.form["stock"]
        costo = request.form["costo"]
        proveedor_id = request.form.get("proveedor_id") or None
        codigo = (insumo.get("codigo_barras") or "").strip()
        if not codigo:
            codigo = generar_codigo_barra("INS")

        foto_archivo = request.files.get("foto")
        reemplaza_foto = foto_archivo and foto_archivo.filename
        quitar_foto = request.form.get("quitar_foto") == "1" and not reemplaza_foto
        nueva_foto_drive = None

        try:
            if reemplaza_foto:
                nueva_foto_drive = subir_foto_drive(foto_archivo, "insumo", id, nombre)
                reconectar_mysql_si_hace_falta()

            cursor.execute("""
                UPDATE insumos
                SET nombre=%s,
                    stock=%s,
                    costo=%s,
                    proveedor_id=%s,
                    codigo_barras=%s
                WHERE id=%s
            """, (nombre, stock, costo, proveedor_id, codigo, id))

            if nueva_foto_drive:
                cursor.execute("""
                    UPDATE insumos
                    SET foto_drive_id=%s,
                        foto_drive_nombre=%s,
                        foto_mime_type=%s,
                        foto_drive_web_view_url=%s,
                        foto_actualizada=NOW()
                    WHERE id=%s
                """, (
                    nueva_foto_drive["foto_drive_id"],
                    nueva_foto_drive["foto_drive_nombre"],
                    nueva_foto_drive["foto_mime_type"],
                    nueva_foto_drive["foto_drive_web_view_url"],
                    id
                ))
            elif quitar_foto:
                cursor.execute("""
                    UPDATE insumos
                    SET foto_drive_id=NULL,
                        foto_drive_nombre=NULL,
                        foto_mime_type=NULL,
                        foto_drive_web_view_url=NULL,
                        foto_actualizada=NULL
                    WHERE id=%s
                """, (id,))

            mysql.connection.commit()
        except (ValueError, RuntimeError, MySQLdb.OperationalError) as error:
            mysql.connection.rollback()
            if nueva_foto_drive:
                eliminar_foto_drive_silencioso(nueva_foto_drive["foto_drive_id"])
            flash(str(error))
            cursor.close()
            return redirect(url_for("editar_insumo", id=id))

        if nueva_foto_drive or quitar_foto:
            eliminar_foto_drive_silencioso(insumo.get("foto_drive_id"))

        cursor.close()
        flash("Insumo actualizado correctamente")
        return redirect(url_for("insumos"))

    cursor.execute("SELECT * FROM proveedores ORDER BY nombre")
    proveedores = cursor.fetchall()
    cursor.close()
    return render_template("editar_insumo.html", insumo=insumo, proveedores=proveedores)


# ELIMINAR INSUMO
# -------------------
@app.route("/eliminar_insumo/<int:id>")
@login_required
@rol_requerido(["admin"])
def eliminar_insumo(id):

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cursor.execute("SELECT foto_drive_id FROM insumos WHERE id=%s", (id,))
    insumo = cursor.fetchone()

    cursor.execute("DELETE FROM insumos WHERE id=%s", (id,))
    mysql.connection.commit()
    cursor.close()

    foto_error = False
    if insumo and insumo.get("foto_drive_id"):
        try:
            eliminar_foto_drive(insumo["foto_drive_id"])
        except RuntimeError:
            foto_error = True

    if foto_error:
        flash("Insumo eliminado, pero no se pudo borrar su foto de Drive")
    else:
        flash("Insumo eliminado correctamente")
    return redirect(url_for("insumos"))

# EXPORTAR INSUMOS A EXCEL
# -------------------
@app.route("/exportar_insumos")
@login_required
@rol_requerido(["admin"])
def exportar_insumos():

    rellenar_codigos_barras_faltantes("insumos")
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("""
        SELECT insumos.*, proveedores.nombre AS proveedor_nombre
        FROM insumos
        LEFT JOIN proveedores ON insumos.proveedor_id = proveedores.id
    """)
    insumos = cursor.fetchall()
    cursor.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Insumos"

    ws.append(["ID", "Nombre", "Stock", "Costo", "Proveedor", "Código Barras"])

    for i in insumos:
        ws.append([
            i["id"],
            i["nombre"],
            i["stock"],
            float(i["costo"]),
            i.get("proveedor_nombre") or i.get("proveedor", ""),
            i.get("codigo_barras", "")
        ])

    file = io.BytesIO()
    wb.save(file)
    file.seek(0)

    return send_file(
        file,
        download_name="insumos.xlsx",
        as_attachment=True
    )

# ETIQUETAS DE INSUMOS EN PDF
# -------------------
@app.route("/etiquetas_insumos_carta")
@login_required
@rol_requerido(["admin"])
def etiquetas_insumos_carta():

    rellenar_codigos_barras_faltantes("insumos")
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)

    elements = []
    styles = getSampleStyleSheet()

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("SELECT nombre, costo, codigo_barras FROM insumos")
    insumos = cursor.fetchall()
    cursor.close()

    etiquetas = []
    fila = []

    for i in insumos:

        contenido = []

        contenido.append(Paragraph(f"<b>{i['nombre']}</b>", styles["Normal"]))
        contenido.append(Spacer(1, 3))
        contenido.append(Paragraph(f"${i['costo']}", styles["Normal"]))
        contenido.append(Spacer(1, 4))

        if i["codigo_barras"]:
            barcode = code128.Code128(
                str(i["codigo_barras"]),
                barHeight=15*mm,
                barWidth=1.0
            )
            contenido.append(barcode)
            contenido.append(Spacer(1, 3))
            contenido.append(Paragraph(str(i["codigo_barras"]), styles["Normal"]))

        etiqueta = Table([[contenido]],
                         colWidths=60*mm,
                         rowHeights=38*mm)

        etiqueta.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.black),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ]))

        fila.append(etiqueta)

        if len(fila) == 3:
            etiquetas.append(fila)
            fila = []

    if fila:
        while len(fila) < 3:
            fila.append("")
        etiquetas.append(fila)

    elements.append(Table(etiquetas, colWidths=[65*mm]*3))

    doc.build(elements)
    buffer.seek(0)

    return send_file(buffer,
                     as_attachment=True,
                     download_name="etiquetas_insumos_carta.pdf",
                     mimetype="application/pdf")

# ETIQUETA DE UN SOLO INSUMO (IMPRESORA TERMICA)
# -------------------
@app.route("/etiqueta_insumo/<int:id>")
@login_required
@rol_requerido(["admin"])
def etiqueta_insumo(id):

    codigo = asegurar_codigo_barra_actual("insumos", id)

    if not codigo:
        flash("Insumo no encontrado")
        return redirect(url_for("insumos"))

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("SELECT nombre, costo, codigo_barras FROM insumos WHERE id=%s", (id,))
    insumo = cursor.fetchone()
    cursor.close()

    return render_template("etiqueta_item_print.html",
                           nombre=insumo["nombre"],
                           precio=insumo["costo"],
                           codigo_barras=insumo["codigo_barras"])


# ETIQUETA PNG DE UN SOLO INSUMO (IMPRESORA TERMICA BLUETOOTH)
# -------------------
@app.route("/etiqueta_insumo/<int:id>/png")
@login_required
@rol_requerido(["admin"])
def etiqueta_insumo_png(id):

    codigo = asegurar_codigo_barra_actual("insumos", id)

    if not codigo:
        flash("Insumo no encontrado")
        return redirect(url_for("insumos"))

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("SELECT nombre, costo, codigo_barras FROM insumos WHERE id=%s", (id,))
    insumo = cursor.fetchone()
    cursor.close()

    imagen = generar_etiqueta_png(insumo["nombre"], insumo["costo"], insumo["codigo_barras"])

    return send_file(
        imagen,
        # Ver comentario equivalente en etiqueta_producto_png sobre por que
        # se usa application/octet-stream en vez de image/png.
        mimetype="application/octet-stream",
        as_attachment=True,
        download_name=f"etiqueta_insumo_{id}.png"
    )


@app.route("/insumo/<int:id>/foto")
@login_required
def foto_insumo(id):
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)
    cursor.execute("""
        SELECT foto_drive_id, foto_drive_nombre, foto_mime_type
        FROM insumos
        WHERE id=%s
    """, (id,))
    insumo = cursor.fetchone()
    cursor.close()

    if not insumo or not insumo.get("foto_drive_id"):
        abort(404)

    try:
        archivo = descargar_foto_drive(insumo["foto_drive_id"])
    except RuntimeError:
        abort(404)

    return send_file(
        archivo,
        mimetype=insumo.get("foto_mime_type") or "image/jpeg",
        download_name=insumo.get("foto_drive_nombre") or f"insumo-{id}.jpg",
        max_age=3600
    )

# -------------------
# PROVEEDORES
# -------------------
@app.route("/proveedores", methods=["GET", "POST"])
@login_required
@rol_requerido(["admin"])
def proveedores():
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    if request.method == "POST":
        nombre = request.form["nombre"]
        direccion_contacto = request.form.get("direccion_contacto")
        telefono = request.form.get("telefono")
        observaciones = request.form.get("observaciones")

        cursor.execute("""
            INSERT INTO proveedores (nombre, direccion_contacto, telefono, observaciones)
            VALUES (%s, %s, %s, %s)
        """, (nombre, direccion_contacto, telefono, observaciones))

        mysql.connection.commit()
        cursor.close()
        flash("Proveedor agregado correctamente")
        return redirect(url_for("proveedores"))

    cursor.execute("""
        SELECT proveedores.*,
               (SELECT COUNT(*) FROM productos WHERE productos.proveedor_id = proveedores.id) AS productos_count,
               (SELECT COUNT(*) FROM insumos WHERE insumos.proveedor_id = proveedores.id) AS insumos_count
        FROM proveedores
        ORDER BY proveedores.nombre
    """)
    proveedores = cursor.fetchall()
    cursor.close()

    return render_template("proveedores.html", proveedores=proveedores)


@app.route("/editar_proveedor/<int:id>", methods=["GET", "POST"])
@login_required
@rol_requerido(["admin"])
def editar_proveedor(id):
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cursor.execute("SELECT * FROM proveedores WHERE id=%s", (id,))
    proveedor = cursor.fetchone()

    if not proveedor:
        cursor.close()
        flash("Proveedor no encontrado")
        return redirect(url_for("proveedores"))

    if request.method == "POST":
        nombre = request.form["nombre"]
        direccion_contacto = request.form.get("direccion_contacto")
        telefono = request.form.get("telefono")
        observaciones = request.form.get("observaciones")

        cursor.execute("""
            UPDATE proveedores
            SET nombre=%s,
                direccion_contacto=%s,
                telefono=%s,
                observaciones=%s
            WHERE id=%s
        """, (nombre, direccion_contacto, telefono, observaciones, id))

        mysql.connection.commit()
        cursor.close()
        flash("Proveedor actualizado correctamente")
        return redirect(url_for("proveedores"))

    cursor.close()
    return render_template("editar_proveedor.html", proveedor=proveedor)


@app.route("/eliminar_proveedor/<int:id>")
@login_required
@rol_requerido(["admin"])
def eliminar_proveedor(id):
    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    cursor.execute("""
        SELECT
            (SELECT COUNT(*) FROM productos WHERE proveedor_id=%s) AS productos_count,
            (SELECT COUNT(*) FROM insumos WHERE proveedor_id=%s) AS insumos_count
    """, (id, id))
    uso = cursor.fetchone()

    if uso and (uso["productos_count"] or uso["insumos_count"]):
        cursor.close()
        flash("No se puede eliminar: el proveedor está ligado a productos o insumos")
        return redirect(url_for("proveedores"))

    cursor.execute("DELETE FROM proveedores WHERE id=%s", (id,))
    mysql.connection.commit()
    cursor.close()

    flash("Proveedor eliminado correctamente")
    return redirect(url_for("proveedores"))


# -------------------
# GASTOS O SERVICIOS A PAGAR
# -------------------
@app.route("/gastos", methods=["GET", "POST"])
@login_required
@rol_requerido(["admin"])
def gastos():

    cursor = mysql.connection.cursor(MySQLdb.cursors.DictCursor)

    if request.method == "POST":

        concepto = request.form["concepto"]
        monto = request.form["monto"]
        fecha = request.form["fecha"]
        descripcion = request.form["descripcion"]

        cursor.execute("""
            INSERT INTO gastos_fijos (concepto, monto, fecha, descripcion)
            VALUES (%s, %s, %s, %s)
        """, (concepto, monto, fecha, descripcion))

        mysql.connection.commit()
        flash("Gasto registrado")
        return redirect(url_for("gastos"))

    cursor.execute("SELECT * FROM gastos_fijos ORDER BY fecha DESC")
    gastos = cursor.fetchall()
    cursor.close()

    return render_template("gastos.html", gastos=gastos)


# -------------------FINAL DE APP-------------------
if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False, use_reloader=False)
